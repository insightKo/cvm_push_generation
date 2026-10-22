# Проверка состава категорий по товарам: свод переносов разборов A–E → Excel на проверку Елене
# и общий файл переносов справочник_CVM_переносы_товаров.json (вход пересборки справочника).
import collections
import json
import sys
from pathlib import Path

import openpyxl

import cvm_category_decisions as decisions
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parent
DST = OUT / 'проверка_состава_категорий_CVM.xlsx'
MOVES = OUT / 'справочник_CVM_переносы_товаров.json'


def main(src_dir):
    src = Path(src_dir)
    tree = {t['category']: t for t in json.load(open(src / 'tree_final.json', encoding='utf-8'))}
    moves, seen = [], set()
    for k in 'ABCDE':
        for m in json.load(open(src / f'moves_{k}.json', encoding='utf-8')):
            if m['product_id'] in seen:      # детские блюда GIPOPO перенесли два разбора — берём первый
                continue
            seen.add(m['product_id'])
            moves.append(m)
    products_by_cat = collections.defaultdict(list)
    for k in 'ABCDE':
        for g, cats in json.load(open(src / f'products_{k}.json', encoding='utf-8')).items():
            for key, items in cats.items():
                products_by_cat[key.split(' ', 1)[1]] += items
    moves, unresolved = decisions.apply(moves, products_by_cat, tree)
    rounds_file = OUT / 'справочник_CVM_проверка_круги.json'
    pending = []
    if rounds_file.exists():
        moves, pending = decisions.apply_rounds(moves, json.load(open(rounds_file, encoding='utf-8')))
        print('вопросов круга 2 без ответа:', len(pending))
    json.dump(moves, open(MOVES, 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    json.dump(list(tree.values()), open(OUT / 'справочник_CVM_дерево_итог.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    if unresolved:
        print('«Здоровое питание» без полки:', unresolved)

    # состав «было» по всем товарам: категория → число товаров
    before = collections.Counter()
    group_of = {name: t['group'] for name, t in tree.items()}
    for k in 'ABCDE':
        for g, cats in json.load(open(src / f'products_{k}.json', encoding='utf-8')).items():
            for key, items in cats.items():
                name = key.split(' ', 1)[1]
                before[name] += len(items)
                group_of.setdefault(name, g)
    out_c, in_c = collections.Counter(m['from'] for m in moves), collections.Counter(m['to'] for m in moves)
    names = sorted(set(before) | set(in_c), key=lambda n: (group_of.get(n, ''), n))

    def status(n):
        after = before[n] - out_c[n] + in_c[n]
        if not before[n] and after:
            return 'новая полка'
        if before[n] and not after:
            return 'убирается'
        return 'меняется' if out_c[n] or in_c[n] else 'без изменений'

    summary = [(group_of.get(n), n, before[n], out_c[n], in_c[n], before[n] - out_c[n] + in_c[n], status(n)) for n in names]

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(ws, head, rows, widths, wrap=False):
        ws.append(head)
        for c in ws[1]:
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical='top')
        for r in rows:
            ws.append(list(r))
        if wrap:
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical='top')
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(head))}{len(rows) + 1}'

    sheet(wb.active, ['Группа', 'Категория', 'Было товаров', 'Ушло', 'Пришло', 'Станет', 'Что с полкой'],
          summary, [22, 40, 12, 9, 9, 10, 16])
    wb.active.title = 'Категории'

    def move_rows(ms):
        ms = sorted(ms, key=lambda m: (group_of.get(m['from'], ''), m['from'], m['to'], m['name'] or ''))
        return [(group_of.get(m['from']), m['from'], m['name'], m['to'], group_of.get(m['to']), m['reason'],
                 'согласованное предложение' if m['source'] == 'предложение' else 'проверка товаров',
                 None if m['sure'] else 'сомнительно', m['product_id']) for m in ms]

    head = ['Группа', 'Из категории', 'Товар', 'В категорию', 'Группа (куда)', 'Почему', 'Основание', 'Уверенность', 'ID товара']
    widths = [20, 34, 56, 34, 20, 46, 22, 12, 11]
    sheet(wb.create_sheet('Сомнительные'), head, move_rows([m for m in moves if not m['sure']]), widths)
    sheet(wb.create_sheet('Найдено проверкой'), head, move_rows([m for m in moves if m['source'] != 'предложение' and m['sure']]), widths)
    sheet(wb.create_sheet('По предложениям'), head, move_rows([m for m in moves if m['source'] == 'предложение']), widths)
    wb.save(DST)
    st = collections.Counter(r[6] for r in summary)
    print(DST.name, 'переносов', len(moves), collections.Counter((m['source'], m['sure']) for m in moves), dict(st))


if __name__ == '__main__':
    main(sys.argv[1])
