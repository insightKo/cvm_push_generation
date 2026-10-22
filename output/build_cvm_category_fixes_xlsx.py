# Переносит правки справочника (справочник_CVM_правки_итог.json от build_cvm_category_fixes.py) в итоговые файлы:
# справочник_категорий_CVM.xlsx (Категории, Группы, Маппинг кодов, Правила разрезов, Товарный маппинг + лист «Правки 19.10»)
# и привязка_товаров_к_категориям_CVM.xlsx (Товар → категория, Сводка по категориям, Несопоставленные).
# Итог файлов = состояние базы после CVM_CATEGORY_FIXES.sql.
import collections
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent
REP = json.loads((OUT / 'справочник_CVM_правки_итог.json').read_text(encoding='utf-8'))
FIX = json.loads((OUT / 'справочник_CVM_правки_замечаний.json').read_text(encoding='utf-8'))
REF = OUT / 'справочник_категорий_CVM.xlsx'
PRD = OUT / 'привязка_товаров_к_категориям_CVM.xlsx'
HOWS = {0: 'несопоставлен', 1: 'код целиком', 2: 'правило слов', 3: 'поштучно'}
METHOD_NAME = {1: 'целиком', 2: 'по словам', 3: 'по товарам'}


def main():
    wb = openpyxl.load_workbook(REF)
    ws = wb['Категории']
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    groups = {r[0].value: r[1].value for r in wb['Группы'].iter_rows(min_row=2) if r[0].value}
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2) if r[0].value}
    cat = {i: {'group': r[col['ID группы']].value, 'name': r[col['Категория']].value, 'unit': r[col['Единица']].value,
               'ns': r[col['Без шкалы']].value} for i, r in rows.items()}
    defs = REP['definitions']
    closed = set(REP['close'])
    final_cat = {int(p): c for p, c in REP['final_category'].items()}
    final_map = {c: v for c, v in REP['final_map'].items()}
    final_man = {int(p): c for p, c in REP['final_manual'].items()}

    # новые полки
    for nc in REP['new_categories']:
        cat[nc['id']] = {'group': nc['group_id'], 'name': nc['name'], 'unit': nc['unit'], 'ns': 'да' if nc['no_scale'] else None}
        ws.append([nc['id'], nc['group_id'], groups[nc['group_id']], nc['name'], nc['unit'], 'да' if nc['no_scale'] else None,
                   'активна', 0, None, None, defs.get(str(nc['id']))])
        rows[nc['id']] = ws[ws.max_row]
    for i, n in REP['renames'].items():
        rows[int(i)][col['Категория']].value = n
        cat[int(i)]['name'] = n
    for i, ns in REP['no_scale'].items():
        rows[int(i)][col['Без шкалы']].value = 'да' if ns else None
    for i in closed:
        rows[i][col['Статус']].value = 'закрыта'
    for i, d in defs.items():
        if i.isdigit() and int(i) in rows and d:
            rows[int(i)][col['Определение']].value = d
    cnt = collections.Counter(final_cat.values())
    whole = collections.defaultdict(list)
    for c, (m, cid) in final_map.items():
        if m == 1:
            whole[cid].append(c)
    for i, r in rows.items():
        r[col['Товаров']].value = cnt.get(i, 0)
        r[col['Ext-коды целиком']].value = ', '.join(sorted(whole.get(i, []))) or None
    g = wb['Группы']
    per_group = collections.Counter(v['group'] for i, v in cat.items() if i not in closed)
    for r in g.iter_rows(min_row=2):
        if r[0].value:
            r[2].value = per_group.get(r[0].value, 0)

    # маппинг кодов
    mp = wb['Маппинг кодов']
    for r in mp.iter_rows(min_row=2):
        c = str(r[0].value) if r[0].value is not None else None
        if c in REP['map_changes']:
            m, cid = final_map[c]
            r[2].value = METHOD_NAME[m]
            r[3].value = cid
            r[4].value = cat[cid]['name'] if cid else None
    # правила разрезов: строки пересобранных кодов заменяются
    rl = wb['Правила разрезов']
    keep = [[c.value for c in r] for r in rl.iter_rows(min_row=2) if str(r[0].value) not in REP['final_rules']]
    for code, rr in REP['final_rules'].items():
        for o, a, n, t in rr:
            keep.append([code, o, (a or '').replace(';', ', ') or None, (n or '').replace(';', ', ') or None, t, cat[t]['name']])
    rl.delete_rows(2, rl.max_row)
    for r in sorted(keep, key=lambda x: (str(x[0]), x[1])):
        rl.append(r)
    # товарный маппинг: новые и изменённые строки
    tm = wb['Товарный маппинг']
    names = {}
    pw = openpyxl.load_workbook(PRD)
    pt = pw['Товар → категория']
    for r in pt.iter_rows(min_row=2, values_only=True):
        names[r[0]] = r[1]
    tmrows = {r[0].value: r for r in tm.iter_rows(min_row=2) if r[0].value}
    moves = REP['moves']
    for p, c in sorted(final_man.items()):
        note = f'правки {FIX["extra"]["date"]}, замечание №{moves[str(p)][2]}' if str(p) in moves else None
        if p in tmrows:
            r = tmrows[p]
            if r[3].value != c:
                r[3].value, r[4].value, r[5].value = c, groups[cat[c]['group']], cat[c]['name']
                r[6].value, r[7].value = 'вручную', note
        else:
            tm.append([p, names.get(p), 'правки замечаний', c, groups[cat[c]['group']], cat[c]['name'], 'вручную', note, None])
    # лист «Правки 19.10»
    if 'Правки 19.10' in wb.sheetnames:
        del wb['Правки 19.10']
    pr = wb.create_sheet('Правки 19.10')
    pr.append(['Замечание', 'ID товара', 'Название', 'Было: категория', 'Стало: категория', 'Стало: группа'])
    for c in pr[1]:
        c.font = Font(bold=True)
    for p, (frm, to, why) in sorted(moves.items(), key=lambda x: (str(x[1][2]).zfill(3), x[1][0], x[1][1])):
        pr.append([why, int(p), names.get(int(p)), cat[frm]['name'], cat[to]['name'], groups[cat[to]['group']]])
    pr.column_dimensions['C'].width = 55
    pr.column_dimensions['D'].width = 35
    pr.column_dimensions['E'].width = 35
    wb.save(REF)

    # привязка товаров
    how = {}
    for r in pt.iter_rows(min_row=2):
        p = r[0].value
        code = str(r[9].value) if r[9].value not in (None, '') else None
        m = final_map.get(code)
        h = 3 if p in final_man else 1 if m and m[0] == 1 else 2 if m and m[0] == 2 else 0
        if final_cat[p] == 1651:
            h = 0
        c = final_cat[p]
        how[p] = h
        r[2].value, r[3].value, r[4].value, r[5].value = groups[cat[c]['group']], cat[c]['name'], c, cat[c]['group']
        r[6].value, r[7].value, r[8].value = cat[c]['unit'], cat[c]['ns'], HOWS[h]
    sv = pw['Сводка по категориям']
    agg = collections.defaultdict(collections.Counter)
    for p, c in final_cat.items():
        agg[c][how[p]] += 1
    sv.delete_rows(2, sv.max_row)
    for c in sorted((i for i in cat if i not in closed), key=lambda i: (groups[cat[i]['group']], cat[i]['name'])):
        a = agg.get(c, collections.Counter())
        sv.append([groups[cat[c]['group']], c, cat[c]['name'], cat[c]['unit'], cat[c]['ns'], sum(a.values()), a[1], a[2], a[3]])
    un = pw['Несопоставленные']
    un.delete_rows(2, un.max_row)
    for r in pt.iter_rows(min_row=2, values_only=True):
        if how[r[0]] == 0:
            un.append(list(r))
    pw.save(PRD)
    print('категорий активных', sum(1 for i, v in cat.items() if i not in closed and i > 1000),
          'несопоставленных товаров', sum(1 for h in how.values() if h == 0), 'HOW', collections.Counter(how.values()))


if __name__ == '__main__':
    main()
