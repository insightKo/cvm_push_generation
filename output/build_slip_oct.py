# -*- coding: utf-8 -*-
"""Купоны на slip — октябрь 2026, все четыре купона в одном документе (решение Елены 30.09.2026):
101393 «Купон 50р. на любую покупку» и 101394–101396 «купон 100р. на 7 дней» (вино / крепкий алкоголь / пиво), выдача 06.10–31.10, прекращается 24.10.
Формулы — как в августе/сентябре (build_slip_sept_50.py, build_slip_sept_alco.py):
  50 ₽:   Доп ТО = (3% − 1%) × клиентов × 459 ₽; скидка = 3% × клиентов × 50; PL = Доп ТО × 30% − скидка; бумага 150 000 ₽ на 6,45 млн клиентов.
  алко:   Доп ТО = (3% − 1%) × 1000 ₽ × клиентов × 54%; скидка = 3% × клиентов × 100 / 1,15; PL = Доп ТО × 30% / 1,15 − скидка; бумага = 10/100 × 0,82 × 54% × клиентов.
Аудитория 50 ₽ — лист «Параметры сегментов» (дамп 07.09.2026): Новые, Активные, Активные LFL без PUSH/APP минус фрод; покупатели алкоголя вычитаются по выборке.
Аудитории алко-купонов — как в расчёте сентября (500 / 500 / 1 250 тыс.), реальные — из выборки. Запуск: python output/build_slip_oct.py"""
import json
from docx import Document
from docx.shared import Pt

NBSP = " "
def rub(x): return f"{int(round(x)):,}".replace(",", NBSP)
mln = lambda x: f"{x / 1e6:.2f}".replace(".", ",")

# --- 50 ₽ на любую покупку ---
d = json.load(open("output/cvm_offline_dump_2026-09-07.json"))
seg = d["Параметры сегментов"]
num = lambda s: int(str(s).replace("\xa0", "").replace(" ", "") or 0)
rows, cur = {}, None
for r in seg:
    v = list(r.values())
    if v[1] in ("Фрод", "Новые", "Активные", "Активные LFL", "Случайные", "Отток", "Спящие", "Аффинитивные"):
        cur = v[1]
        if v[3]: rows[(cur, "all")] = (num(v[3]), num(v[4]))
    elif v[1] in ("omni", "offline", "e-commerce") and cur:
        rows[(cur, v[1])] = (num(v[3]), num(v[4]))
aud = sum(rows.get((s, ch), (0, 0))[0] - rows.get((s, ch), (0, 0))[1] for s in ("Новые", "Активные", "Активные LFL") for ch in ("offline", "omni"))
clients50 = aud - (rows[("Фрод", "all")][0] - rows[("Фрод", "all")][1])
dop50 = 0.02 * clients50 * 459; disc50 = 0.03 * clients50 * 50; pl50 = dop50 * 0.30 - disc50; paper50 = clients50 * 150_000 / 6_450_000

# --- 100 ₽ на категорию (алкоголь) ---
CHECK, COEF, MARGIN, VAT, COUPON, RESP, CANN = 1000, 0.54, 0.30, 1.15, 100, 0.03, 0.01
alco = [(101394, "Вина купон на 100р. на 7 дней", "вино и игристое", 519, 500_000),
        (101395, "Крепкий алкоголь купон на 100р. на 7 дней", "крепкий алкоголь", 479, 500_000),
        (101396, "Пиво купон на 100р. на 7 дней", "пиво", 105, 1_250_000)]
table = [(101393, "Купон 50р. на любую покупку", "Активные, Новые без PUSH/APP, без покупок алкоголя", "любая покупка", clients50, "50 ₽", dop50, disc50, pl50, paper50)]
for n, name, cat, avg, cl in alco:
    dop = (RESP - CANN) * CHECK * cl * COEF; disc = RESP * cl * COUPON / VAT; pl = dop * MARGIN / VAT - disc; comm = 0.10 * 0.82 * COEF * cl
    table.append((n, name, "Активные, Новые — покупатели категории за 8 недель, без PUSH/APP", cat, cl, f"100 ₽ на чек от 1000 ₽ (≈{round(COUPON / avg * 100)}% от закупки)", dop, disc, pl, comm))
sum_to = sum(t[6] for t in table); sum_disc = sum(t[7] for t in table); sum_pl = sum(t[8] for t in table); sum_comm = sum(t[9] for t in table)

doc = Document(); st = doc.styles["Normal"].font; st.name = "Calibri"; st.size = Pt(11)
def p(text, bold=False, size=11):
    par = doc.add_paragraph(); r = par.add_run(text); r.bold = bold; r.font.size = Pt(size); return par
p("Акции с использованием канала: слип-чек. 06.10-31.10.2026", bold=True, size=14)
p("Целевая аудитория: офлайн-клиенты ДИКСИ, для которых слип-чек — основной канал персональной коммуникации (без приложения и PUSH). "
  "Покупатели алкоголя получают купон 100 ₽ на свою категорию, остальные Активные и Новые — купон 50 ₽ на любую покупку.")
p("Цели:")
for g in ["дополнительный трафик по Активным и Новым (купон 50р. на любую покупку)",
          "дополнительный трафик и частота по покупателям вина, крепкого алкоголя и пива (купон 100р. на чек от 1000р.)",
          "возврат за повторной покупкой в течение 7 дней после выдачи слип-чека"]:
    doc.add_paragraph(g, style="List Bullet")
p("Даты проведения акций: 06.10–31.10.2026. Выдача купонов прекращается 24.10.2026 — купон действует 7 дней.")
p("Механика 50р.: за период акции после любой покупки клиенту выдаётся слип-чек с QR-купоном на 50р. (применяется при предъявлении купона и карты на 1 покупку, "
  "не более 2 применений в сутки; не действует на социально значимые товары, алкоголь, табак, лотерейные билеты и промотовары по жёлтым ценникам).")
p("Механика 100р.: после покупки товаров категории клиенту выдаётся слип-чек с купоном на 100р. Скидка предоставляется при сумме чека от 1000р. и распределяется по всем товарам чека; "
  "для клиента — скидка на вино / крепкий алкоголь / пиво, условие о сумме чека — мелким шрифтом на купоне. Не суммируется с другими акциями, не действует на промотовары по жёлтым ценникам.")
p(f"Затраты на коммуникацию на чеках: {round(sum_comm / 1000)} тыс ₽")
p("Планируемый эффект:")
p(f"+{mln(sum_to)} млн ₽ (Доп ТО); PL +{mln(sum_pl)} млн ₽", bold=True)
doc.add_paragraph()
headers = ["НОМЕР", "Название промо", "Старт", "Финиш", "Сегмент МС", "Категория", "Клиентов", "отклик, %", "Каннибал., %", "Скидка (купон)", "Доп ТО (план), р.", "Скидка, р.", "PL, р.", "Расход бумаги, р."]
t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
for i, h in enumerate(headers):
    c = t.rows[0].cells[i].paragraphs[0].add_run(h); c.bold = True; c.font.size = Pt(8)
for n, name, segm, cat, cl, coupon, dop, disc, pl, comm in table:
    cells = t.add_row().cells
    for i, v in enumerate([str(n), name, "06.10.", "31.10.", segm, cat, rub(cl), "3%", "1%", coupon, rub(dop), rub(disc), rub(pl), rub(comm)]):
        run = cells[i].paragraphs[0].add_run(v); run.font.size = Pt(8)
cells = t.add_row().cells
for i, v in enumerate(["", "ИТОГО", "", "", "", "", rub(sum(x[4] for x in table)), "", "", "", rub(sum_to), rub(sum_disc), rub(sum_pl), rub(sum_comm)]):
    run = cells[i].paragraphs[0].add_run(v); run.bold = True; run.font.size = Pt(8)
doc.add_paragraph()
p("Численность 50р. — по листу «Параметры сегментов» до вычета покупателей алкоголя (доля определится выборкой); численности алко-купонов — допущения расчёта сентября, реальные — из выборки.", size=9)
out = "output/Купоны на slip — 06.10-31.10.docx"
doc.save(out); print("SAVED:", out)
for n, name, segm, cat, cl, coupon, dop, disc, pl, comm in table:
    print(f"{n} {name[:42]:42} клиентов {rub(cl):>10}  ДопТО {rub(dop):>11}  скидка {rub(disc):>10}  PL {rub(pl):>10}  бумага {rub(comm):>7}")
print(f"ИТОГО  ДопТО {rub(sum_to)}  скидка {rub(sum_disc)}  PL {rub(sum_pl)}  бумага {rub(sum_comm)}")
