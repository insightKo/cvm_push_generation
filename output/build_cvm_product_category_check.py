# Привязка «товар → наша категория» по всем товарам I_PRODUCT — для проверки глазами.
# Логика та же, что в dbo.cvm_category_refresh: поштучно → код целиком → правило слов → «Несопоставленные».
# Вход: справочник_категорий_CVM.xlsx и выгрузка I_PRODUCT (csv: P,ID,NAME,EXT / E,EXT,NAME,).
import collections
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

import build_cvm_category_reference as ref

OUT = Path(__file__).resolve().parent
SRC = OUT / 'справочник_категорий_CVM.xlsx'
DST = OUT / 'привязка_товаров_к_категориям_CVM.xlsx'
HOW = {3: 'поштучно', 1: 'код целиком', 2: 'правило слов', 0: 'несопоставлен'}


def main(products_csv):
    wb = openpyxl.load_workbook(SRC, read_only=True)
    rd = lambda s: list(wb[s].iter_rows(values_only=True))[1:]
    cats = {r[0]: r for r in rd('Категории')}
    by_key = {(r[2], r[3]): r[0] for r in cats.values()}
    codes = {r[0]: r for r in rd('Маппинг кодов')}
    manual = {r[0]: r[3] for r in rd('Товарный маппинг')}
    rules = collections.defaultdict(list)
    for r in rd('Правила разрезов'):
        rules[str(r[0])].append(r)
    unmapped = next(k for k, c in cats.items() if c[3] == 'Несопоставленные')

    prods, ext_name = {}, {}
    with open(products_csv, encoding='utf-8') as f:
        for row in csv.reader(f):
            if row[0] == 'P':
                pid = int(row[1])
                old = prods.get(pid, ('', ''))
                prods[pid] = (max(old[0], row[2]), max(old[1], row[3]))
            elif row[0] == 'E':
                ext_name.setdefault(row[1], row[2])

    rows = []
    for pid, (name, code) in prods.items():
        m = codes.get(code)
        if pid in manual:
            cat, how = manual[pid], 3
        elif m and m[2] == 'целиком':
            cat, how = m[3], 1
        elif m and m[2] == 'по словам':
            cat, how = unmapped, 2
            for _, _, any_w, none_w, rid, _ in sorted(rules.get(code, []), key=lambda x: x[1]):
                inc = [] if str(any_w).startswith('—') else [w.strip() for w in str(any_w).split(',')]
                exc = [w.strip() for w in str(none_w).split(',')] if none_w else []
                if exc and ref.word_hit(name, exc):
                    continue
                if not inc or ref.word_hit(name, inc):
                    cat = rid
                    break
        else:
            cat, how = unmapped, 0
        if cat == unmapped:
            how = 0
        c = cats[cat]
        rows.append((pid, name, c[2], c[3], cat, c[1], c[4], c[5], HOW[how],
                     code or None, (m[1] if m else None) or ext_name.get(code)))
    rows.sort(key=lambda r: (r[2], r[3], r[1]))

    out = openpyxl.Workbook(write_only=True)
    bold = Font(bold=True)

    def sheet(title, head, data, widths):
        ws = out.create_sheet(title)
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = 'C2' if title != 'Таблица в базе' else 'A2'
        ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(head))}{len(data) + 1}'
        cells = []
        for h in head:
            cell = openpyxl.cell.WriteOnlyCell(ws, value=h)
            cell.font = bold
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cells.append(cell)
        ws.append(cells)
        for r in data:
            ws.append(list(r))

    cnt = collections.defaultdict(collections.Counter)
    for r in rows:
        cnt[r[4]][r[8]] += 1
        cnt[r[4]]['всего'] += 1
    summary = sorted(((c[2], c[0], c[3], c[4], c[5], cnt[k]['всего'], cnt[k]['код целиком'],
                       cnt[k]['правило слов'], cnt[k]['поштучно']) for k, c in cats.items()),
                     key=lambda r: (r[0], r[2]))
    sheet('Сводка по категориям', ['Наша CVM группа', 'ID нашей CVM категории', 'Наша CVM категория', 'Единица полки', 'Без шкалы', 'Товаров',
                                   'по коду целиком', 'по правилу слов', 'поштучно'],
          summary, [22, 12, 36, 9, 10, 10, 13, 13, 11])
    head = ['ID товара (ID_PRODUCT)', 'Название товара', 'Наша CVM группа', 'Наша CVM категория',
            'ID нашей CVM категории (ID_CVM_CATEGORY)', 'ID нашей CVM группы (ID_CVM_GROUP)', 'Единица полки', 'Без шкалы',
            'Как отнесён (HOW)', 'Код категории ДИКСИ (ID_CATEGORY_5_ext)', 'Название категории ДИКСИ']
    widths = [12, 58, 22, 34, 14, 12, 9, 9, 15, 15, 44]
    fields = [
        ('Таблица', 'dbo.I_CVM_PRODUCT_CATEGORY', '', 'Итог «товар → наша категория» по всем товарам I_PRODUCT (ID_COMPANY = 1). '
         'Её читают все витрины, отчёты и модели. Заполняет ежедневная процедура dbo.cvm_category_refresh '
         '(шаг CVM_CATEGORY в «MCI Regular update»). Одна строка = один товар. Товары, ушедшие из I_PRODUCT, остаются.'),
        ('Поле', 'ID_PRODUCT', 'bigint, ключ', 'ID товара — как в I_PRODUCT.'),
        ('Поле', 'ID_CVM_CATEGORY', 'int', 'Наша категория — ID из dbo.I_CVM_CATEGORY (1001 и дальше).'),
        ('Поле', 'ID_CVM_GROUP', 'int', 'Наша группа категории — ID из dbo.I_CVM_CATEGORY (1–20).'),
        ('Поле', 'ID_CATEGORY_5_ext', 'nvarchar(50)', 'Код категории ДИКСИ у товара (вход маппинга, для анализа не использовать).'),
        ('Поле', 'HOW', 'tinyint', 'Как отнесён товар: 1 — код ДИКСИ целиком; 2 — правило по словам названия; '
         '3 — поштучная привязка; 0 — «Несопоставленные» (в ценовую шкалу не входит). '
         'На листе «Товар → категория» — словами в колонке «Как отнесён».'),
        ('Поле', 'DATE_UPDATE', 'date', 'Дата прогона, в который строка добавлена или изменилась (заполнится при запуске).'),
    ]
    sheet('Таблица в базе', ['Что', 'Название', 'Тип', 'Значение'], fields, [10, 30, 16, 110])
    sheet('Товар → категория', head, rows, widths)
    marks = [w for inc, _, _ in ref.WATER for w in inc]
    sheet('Вода без пометки газа', head,
          [r for r in rows if r[3] in ('Вода газированная', 'Вода негазированная') and r[8] == 'правило слов'
           and not ref.word_hit(r[1], marks)], widths)
    sheet('Несопоставленные', head, [r for r in rows if r[8] == 'несопоставлен'], widths)
    out.save(DST)
    print(DST.name, len(rows), collections.Counter(r[8] for r in rows))


if __name__ == '__main__':
    main(sys.argv[1])
