"""Тесты export.py без баз: подставные источники портала и mci_model, хранилище во временном каталоге.

Запуск на ферме: python3 -m pytest -q tools/campaign-lists/test_export.py
"""
import datetime as dt
import hashlib
import json
import os
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export as ex  # noqa: E402

NOW = dt.datetime(2026, 10, 8, 14, 20, tzinfo=ex.MSK)
TODAY = NOW.date()


class Portal:
    def __init__(self, promos, podborki=(), coupons=()):
        self.promos = promos
        self.podborki_rows = list(podborki)
        # список строк cvm.coupon_publish; None — представления нет (миграция портала не применена); Exception — ошибка
        self.coupon_rows = list(coupons) if isinstance(coupons, (list, tuple)) else coupons

    def approved_promos(self, start_from):
        return [p for p in self.promos if p.start_date >= start_from]

    def podborki(self):
        if isinstance(self.podborki_rows, Exception):
            raise self.podborki_rows
        return self.podborki_rows

    def coupons(self):
        if isinstance(self.coupon_rows, Exception):
            raise self.coupon_rows
        return self.coupon_rows


class Mci:
    """counts — что вернёт каждый следующий подсчёт (последнее значение повторяется); guids — список акции;
    portal — номера акций, у которых строки с LOAD_TO_ML = 2 (список портала)."""

    def __init__(self, guids, counts=None, portal=()):
        self.guids = guids
        self.counts = list(counts) if counts is not None else [len(guids)]
        self.count_calls = 0
        self.portal = set(portal)

    def offer_count(self, promo):
        value = self.counts[min(self.count_calls, len(self.counts) - 1)]
        self.count_calls += 1
        return value, (value if promo in self.portal else 0)

    def offer_guids(self, promo):
        return iter(self.guids)


class Sleep:
    def __init__(self):
        self.calls = []

    def __call__(self, seconds):
        self.calls.append(seconds)


def promo(days=2, number=101500):
    return ex.Promo(number=number, title="Тест", start_date=TODAY + dt.timedelta(days=days), approved_by="elena")


def cfg(tmp_path):
    return ex.Config(storage=tmp_path, state=tmp_path / "_state", retention_days=14,
                     stable_seconds=300, stable_tries=6)


def guids(n):
    return [str(uuid.UUID(int=i + 1)).upper() for i in range(n)]


def events(tmp_path):
    p = tmp_path / "_state" / "journal.jsonl"
    return [json.loads(line)["event"] for line in p.read_text(encoding="utf-8").splitlines()] if p.exists() else []


def csv_of(tmp_path, folder="inbox"):
    return next((tmp_path / folder).glob("*.csv"))


def test_stable_list_exported_once(tmp_path):
    sleep, mci = Sleep(), Mci(guids(5))
    assert ex.run(cfg(tmp_path), Portal([promo()]), mci, NOW, sleep) == 0
    assert sleep.calls == [300]  # подсчёт, 5 минут, подсчёт — совпало
    csv = csv_of(tmp_path)
    assert csv.name == f"101500_{TODAY + dt.timedelta(days=2):%Y%m%d}.csv"
    assert csv.read_bytes().decode().split("\r\n")[:-1] == guids(5)
    manifest = json.loads(csv.with_name(csv.stem + ".manifest.json").read_text(encoding="utf-8"))
    assert manifest["count"] == 5 and manifest["promo"] == 101500 and len(manifest["sha256"]) == 64
    assert oct(csv.stat().st_mode & 0o777) == "0o600"
    assert events(tmp_path) == ["exported"]
    # Следующий прогон ничего не делает: акция уже выгружена.
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(7)), NOW + dt.timedelta(hours=1), Sleep()) == 0
    assert events(tmp_path) == ["exported"]


def test_loading_list_waits_until_stable(tmp_path):
    sleep, mci = Sleep(), Mci(guids(8), counts=[3, 6, 8, 8])
    assert ex.run(cfg(tmp_path), Portal([promo()]), mci, NOW, sleep) == 0
    assert sleep.calls == [300, 300, 300]
    assert csv_of(tmp_path).read_bytes().count(b"\r\n") == 8


def test_still_loading_after_all_tries_is_postponed(tmp_path):
    sleep = Sleep()
    mci = Mci(guids(9), counts=[1, 2, 3, 4, 5, 6, 7, 8, 9])
    assert ex.run(cfg(tmp_path), Portal([promo()]), mci, NOW, sleep) == 0
    assert len(sleep.calls) == 6
    assert not list((tmp_path / "inbox").iterdir())
    assert events(tmp_path) == ["not_ready"]


def test_empty_list_not_exported(tmp_path):
    sleep = Sleep()
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci([], counts=[0]), NOW, sleep) == 0
    assert sleep.calls == []
    assert not list((tmp_path / "inbox").iterdir())


def test_window_not_started_any_distance(tmp_path):
    """Выгружается сразу после согласования: и через 30 дней, и сегодня; уже начавшаяся — нет."""
    portal = Portal([promo(days=30, number=1), promo(days=0, number=2), promo(days=-1, number=3)])
    ex.run(cfg(tmp_path), portal, Mci(guids(2)), NOW, Sleep())
    assert sorted(p.name.split("_")[0] for p in (tmp_path / "inbox").glob("*.csv")) == ["1", "2"]

def test_count_mismatch_leaves_nothing(tmp_path):
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(5), counts=[6]), NOW, Sleep()) == 1
    assert not list((tmp_path / "inbox").iterdir())
    assert events(tmp_path) == ["export_failed"]
    # Не отмечена как выгруженная — следующий прогон попробует снова.
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(5)), NOW + dt.timedelta(hours=1), Sleep()) == 0
    assert events(tmp_path)[-1] == "exported"


def test_null_guid_written_as_empty_line(tmp_path):
    ex.run(cfg(tmp_path), Portal([promo()]), Mci([guids(1)[0], None]), NOW, Sleep())
    assert csv_of(tmp_path).read_bytes().decode().split("\r\n") == [guids(1)[0], "", ""]


def test_journal_and_state_have_no_guids(tmp_path):
    ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(5)), NOW, Sleep())
    text = (tmp_path / "_state" / "journal.jsonl").read_text(encoding="utf-8")
    for state in (tmp_path / "_state" / "state").iterdir():
        text += state.read_text(encoding="utf-8")
    assert not any(g in text for g in guids(5))


def test_sftp_side_has_no_journal_or_state(tmp_path):
    ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(3)), NOW, Sleep())
    visible = {p.name for p in tmp_path.iterdir()} - {"_state"}
    # podborki/ — файлы подборок товаров под пуши (с 20.10.2026), coupons/ — картинки купонов (с 21.10.2026),
    # клиентов в них нет.
    assert visible == {"inbox", "processing", "sent", "failed", "receipts", "podborki", "coupons"}


def test_retention(tmp_path):
    storage = ex.Storage(tmp_path, tmp_path / "_state")
    old = storage.path("sent") / "1_20260901.csv"
    fresh = storage.path("failed") / "2_20261001.csv"
    old.write_text("x"), fresh.write_text("x")
    ts = (NOW - dt.timedelta(days=15)).timestamp()
    os.utime(old, (ts, ts))
    ex.cleanup(storage, 14, NOW)
    assert not old.exists() and fresh.exists()
    assert events(tmp_path) == ["deleted"]


def test_many_promos_share_one_wait(tmp_path):
    """Десяток акций в окне ждёт одну общую паузу, а не по 5 минут на каждую: иначе прогон не уложится в час."""
    sleep = Sleep()
    portal = Portal([promo(number=n) for n in range(1, 13)])
    assert ex.run(cfg(tmp_path), portal, Mci(guids(4)), NOW, sleep) == 0
    assert sleep.calls == [300]
    assert len(list((tmp_path / "inbox").glob("*.csv"))) == 12


class BrokenCount(Mci):
    def offer_count(self, promo):
        if promo == 1:
            raise RuntimeError("timeout")
        return super().offer_count(promo)


def test_count_error_of_one_promo_does_not_stop_others(tmp_path):
    portal = Portal([promo(number=1), promo(number=2)])
    assert ex.run(cfg(tmp_path), portal, BrokenCount(guids(3)), NOW, Sleep()) == 1
    assert [p.name.split("_")[0] for p in (tmp_path / "inbox").glob("*.csv")] == ["2"]
    assert sorted(events(tmp_path)) == ["export_failed", "exported"]


class BrokenStream(Mci):
    def offer_guids(self, promo):
        yield self.guids[0]
        raise RuntimeError("connection reset")


def test_stream_error_leaves_no_part_file(tmp_path):
    assert ex.run(cfg(tmp_path), Portal([promo()]), BrokenStream(guids(3)), NOW, Sleep()) == 1
    assert not list((tmp_path / "inbox").iterdir())


def test_no_duplicate_when_mark_lost(tmp_path):
    """Прогон оборвался после записи файла, но до отметки: следующий прогон не выгружает акцию второй раз."""
    ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(3)), NOW, Sleep())
    base = f"101500_{TODAY + dt.timedelta(days=2):%Y%m%d}"
    (tmp_path / "_state" / "state" / f"{base}.exported").unlink()
    for p in (tmp_path / "inbox").iterdir():  # партнёр уже забрал файл в работу
        p.replace(tmp_path / "processing" / p.name)
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(3)), NOW + dt.timedelta(hours=1), Sleep()) == 0
    assert not list((tmp_path / "inbox").iterdir())
    assert events(tmp_path).count("exported") == 1


def test_mark_kept_until_start_no_second_export(tmp_path):
    """Согласована за 30 дней до старта: через 15 дней файл из sent/ удалён сроком хранения, а отметка — нет,
    иначе акция выгрузилась бы второй раз."""
    p = promo(days=30)
    ex.run(cfg(tmp_path), Portal([p]), Mci(guids(3)), NOW, Sleep())
    for f in (tmp_path / "inbox").iterdir():
        f.replace(tmp_path / "sent" / f.name)
    ts = (NOW - dt.timedelta(days=15)).timestamp()
    for f in list((tmp_path / "sent").iterdir()) + list((tmp_path / "_state" / "state").iterdir()):
        os.utime(f, (ts, ts))
    later = NOW + dt.timedelta(days=15)
    assert ex.run(cfg(tmp_path), Portal([p]), Mci(guids(3)), later, Sleep()) == 0
    assert not list((tmp_path / "sent").iterdir())
    assert events(tmp_path).count("exported") == 1


def test_mark_removed_after_start_plus_retention(tmp_path):
    storage = ex.Storage(tmp_path, tmp_path / "_state")
    old = storage.exported_mark(f"1_{TODAY - dt.timedelta(days=15):%Y%m%d}")
    soon = storage.exported_mark(f"2_{TODAY + dt.timedelta(days=30):%Y%m%d}")
    for m in (old, soon):
        m.write_text("x")
        ts = (NOW - dt.timedelta(days=20)).timestamp()
        os.utime(m, (ts, ts))
    ex.cleanup(storage, 14, NOW)
    assert not old.exists() and soon.exists()


def test_deadline_counts_from_run_start(tmp_path, monkeypatch):
    """Прод: одна пауза, срок 10 минут с начала подсчёта. Первые подсчёты шли 6 минут — пауза не начинается."""
    clock = [0.0]
    monkeypatch.setattr(ex.time, "monotonic", lambda: clock[0])

    class Slow(Mci):
        def offer_count(self, n):
            clock[0] += 360
            return super().offer_count(n)

    sleep = Sleep()
    prod = ex.Config(storage=tmp_path, state=tmp_path / "_state", retention_days=14)
    assert ex.run(prod, Portal([promo()]), Slow(guids(3)), NOW, sleep) == 0
    assert sleep.calls == [] and events(tmp_path) == ["not_ready"]


def test_prod_defaults_one_pause_then_next_run(tmp_path):
    """Прод: одна пауза 5 минут; число изменилось — акция ждёт следующего прогона, без второй паузы."""
    sleep = Sleep()
    prod = ex.Config(storage=tmp_path, state=tmp_path / "_state", retention_days=14)
    assert ex.run(prod, Portal([promo()]), Mci(guids(4), counts=[3, 4]), NOW, sleep) == 0
    assert sleep.calls == [300] and events(tmp_path) == ["not_ready"]
    assert ex.run(prod, Portal([promo()]), Mci(guids(4)), NOW + dt.timedelta(minutes=15), Sleep()) == 0
    assert events(tmp_path)[-1] == "exported"


# ---------- списки портала (решения Елены 10.10.2026) ----------

def test_portal_list_in_mci_is_not_exported(tmp_path):
    """LOAD_TO_ML = 2 — список портала: файл в обмен передал портал, по базе задание его не выгружает и не ждёт паузу."""
    sleep = Sleep()
    portal = Portal([promo(number=500001), promo(number=101500)])
    assert ex.run(cfg(tmp_path), portal, Mci(guids(4), portal={500001}), NOW, sleep) == 0
    assert [p.name.split("_")[0] for p in (tmp_path / "inbox").glob("*.csv")] == ["101500"]
    assert sorted(events(tmp_path)) == ["exported", "skipped_portal_list"]
    assert ex.run(cfg(tmp_path), portal, Mci(guids(4), portal={500001}), NOW + dt.timedelta(minutes=15), Sleep()) == 0
    assert events(tmp_path).count("skipped_portal_list") == 1                        # отмечена — больше не проверяется


def outbox_pair(outbox, number=500000, n=3, start=None, **over):
    start = start or TODAY + dt.timedelta(days=9)
    base = f"{number}_{start:%Y%m%d}"
    body = "".join(g + "\r\n" for g in guids(n)).encode()
    m = {"promo": number, "title": "Омни · пиво", "start_date": str(start), "approved_by": "elena", "count": n,
         "sha256": hashlib.sha256(body).hexdigest(), "file": f"{base}.csv", "exported_at": "2026-10-10T20:00:00+03:00", **over}
    outbox.mkdir(exist_ok=True)
    (outbox / f"{base}.csv").write_bytes(body)
    (outbox / f"{base}.manifest.json").write_text(json.dumps(m, ensure_ascii=False), encoding="utf-8")
    return base, body


def portal_cfg(tmp_path):
    return ex.Config(storage=tmp_path / "lists", state=tmp_path / "_state", retention_days=14,
                     portal_outbox=tmp_path / "outbox")


def test_portal_pair_moves_to_inbox(tmp_path):
    base, body = outbox_pair(tmp_path / "outbox")
    assert ex.run(portal_cfg(tmp_path), Portal([]), Mci([]), NOW, Sleep()) == 0
    inbox = tmp_path / "lists" / "inbox"
    assert (inbox / f"{base}.csv").read_bytes() == body
    m = json.loads((inbox / f"{base}.manifest.json").read_text(encoding="utf-8"))
    assert m["promo"] == 500000 and m["count"] == 3 and m["sha256"] == hashlib.sha256(body).hexdigest()
    assert m["title"] == "Омни · пиво" and m["file"] == f"{base}.csv"
    assert not list((tmp_path / "outbox").iterdir())                                   # пара забрана — портал увидит «на SFTP»
    assert events(tmp_path) == ["exported"]
    journal = (tmp_path / "_state" / "journal.jsonl").read_text(encoding="utf-8")
    assert '"source": "portal"' in journal and not any(g in journal for g in guids(3))
    # та же пара ещё раз (повтор создания в портале) — второй выгрузки нет
    outbox_pair(tmp_path / "outbox")
    assert ex.run(portal_cfg(tmp_path), Portal([]), Mci([]), NOW, Sleep()) == 0
    assert events(tmp_path) == ["exported", "portal_duplicate"] and not list((tmp_path / "outbox").iterdir())


def test_portal_pair_with_bad_content_stays(tmp_path):
    """Подмена после паспорта, строка не GUID, чужой номер в паспорте — в inbox ничего, пара на месте, причина в журнале."""
    outbox = tmp_path / "outbox"
    base, body = outbox_pair(outbox, number=500000)
    (outbox / f"{base}.csv").write_bytes(body.replace(b"0000-0000", b"0000-1111", 1))     # SHA-256 не совпал
    base2, _ = outbox_pair(outbox, number=500001)
    (outbox / f"{base2}.csv").write_bytes(b"DROP TABLE x\r\n" * 3)                      # не CRM_GUID
    outbox_pair(outbox, number=500002, promo=500099)                                      # номер в паспорте не из имени
    assert ex.run(portal_cfg(tmp_path), Portal([]), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "lists" / "inbox").iterdir())
    assert len(list(outbox.glob("*.manifest.json"))) == 3
    assert events(tmp_path) == ["export_failed"] * 3


def test_portal_symlink_is_not_followed(tmp_path):
    """В папку пишет API портала: ссылка на чужой файл (например, список другой акции) в обмен не попадает."""
    outbox = tmp_path / "outbox"
    base, body = outbox_pair(outbox)
    secret = tmp_path / "other.csv"
    secret.write_bytes(body)
    (outbox / f"{base}.csv").unlink()
    (outbox / f"{base}.csv").symlink_to(secret)
    assert ex.run(portal_cfg(tmp_path), Portal([]), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "lists" / "inbox").iterdir()) and events(tmp_path) == ["export_failed"]


def test_portal_pair_without_manifest_waits(tmp_path):
    """Паспорт появляется последним: без него список ещё пишется — не трогаем."""
    outbox = tmp_path / "outbox"
    base, _ = outbox_pair(outbox)
    (outbox / f"{base}.manifest.json").unlink()
    assert ex.run(portal_cfg(tmp_path), Portal([]), Mci([]), NOW, Sleep()) == 0
    assert (outbox / f"{base}.csv").exists() and events(tmp_path) == []


def test_campaign_title_up_to_50_and_not_ml_name():
    """Имя кампании «<номер>_<title>» — название словами целиком до 50 знаков и не совпадает с PROMO_NAME_ML, иначе
    Manzana его не примет (решение Елены 19.10.2026)."""
    p = ex.Promo(number=101439, title="Активируй 50% кешбэка на выпечку собственного производства",
                 start_date=dt.date(2026, 10, 15), approved_by="CVM offline")

    class WithName:
        def __init__(self, name):
            self.name = name

        def promo_name_ml(self, n):
            return self.name

    title = ex.campaign_title(p, WithName("101439_50 кешбэка выпечку"))
    assert title == "Активируй 50% кешбэка на выпечку"
    assert len(f"{p.number}_{title}") <= 50
    assert ex.campaign_title(p, object()) == title

    short = ex.Promo(number=101425, title="Баланс баллов", start_date=dt.date(2026, 10, 15), approved_by="CVM offline")
    assert ex.campaign_title(short, WithName("101425_Баланс баллов")) == "Баланс баллов PUSH"
    assert ex.campaign_title(short, WithName("101425_баланс БАЛЛОВ")) == "Баланс баллов PUSH"
    assert ex.campaign_title(short, WithName(None)) == "Баланс баллов"

    # у акции нет названия в портале — слова из имени для Online, и всё равно не то же имя
    untitled = ex.Promo(number=101425, title=None, start_date=dt.date(2026, 10, 15), approved_by="CVM offline")
    assert ex.campaign_title(untitled, WithName("101425_Баланс баллов")) == "Баланс баллов PUSH"

    one = ex.Promo(number=101426, title="Х" * 80 + "\nвторое", start_date=dt.date(2026, 10, 15), approved_by="CVM offline")
    assert ex.campaign_title(one, object()) == "Х" * 43


# ---------- подборки (решение Елены 20.10.2026) ----------

def podborka(day, file="подборка_детская_101397.xlsx", number=101397, episode="Серия 3 — Сок", rows=None):
    return {"file_name": file, "promo_number": number, "series_title": "Подборка детская",
            "send_date": dt.date(2026, 10, day), "episode": episode,
            "rows": rows if rows is not None else [{"block": "Соки", "sku": "2000536279", "name": "СОК ФРУТОНЯНЯ 0,2Л",
                                                    "cat5": 61050101}]}


def sheet(path):
    import openpyxl
    ws = openpyxl.load_workbook(path).active
    return ws.title, [[c.value for c in r] for r in ws.iter_rows()], ws


def test_podborka_file_written_once_with_all_dates(tmp_path):
    folder = tmp_path / "podborki"
    portal = Portal([], [podborka(27, episode="Серия 4 — Печенье"), podborka(20)])
    assert ex.run(cfg(tmp_path), portal, Mci([]), NOW, Sleep()) == 0
    path = folder / "подборка_детская_101397.xlsx"
    title, rows, ws = sheet(path)
    assert title == "SKU по акциям октябрь"
    assert rows[0] == list(ex.PODBORKA_HEAD) and ws["A1"].font.b
    # Даты по порядку: 20.10, потом 27.10 — следующая согласованная коммуникация дописывает свои строки в тот же файл.
    assert [r[2] for r in rows[1:]] == ["20.10.2026", "27.10.2026"]
    assert rows[1] == [101397, "Подборка детская", "20.10.2026", "Серия 3 — Сок", "Соки", "СОК ФРУТОНЯНЯ 0,2Л", "2000536279"]
    assert oct(path.stat().st_mode & 0o777) == "0o600"
    assert events(tmp_path) == ["podborka_written"]
    # Состав не менялся — файл не переписывается.
    ex.run(cfg(tmp_path), portal, Mci([]), NOW, Sleep())
    assert events(tmp_path) == ["podborka_written"]
    # Файл пропал с диска — пишется снова.
    path.unlink()
    ex.run(cfg(tmp_path), portal, Mci([]), NOW, Sleep())
    assert events(tmp_path) == ["podborka_written", "podborka_written"] and path.exists()


def test_podborka_withdrawn_file_removed_manual_file_kept(tmp_path):
    folder = tmp_path / "podborki"
    folder.mkdir()
    manual = folder / "подборки_октябрь_7.xlsx"
    manual.write_bytes(b"x")
    ex.run(cfg(tmp_path), Portal([], [podborka(20), podborka(23, file="подборка_пиво_101402.xlsx", number=101402)]),
           Mci([]), NOW, Sleep())
    assert (folder / "подборка_пиво_101402.xlsx").exists()
    # Пуш пива больше не согласован — его подборки нет в представлении, файл уходит; у детской сменился состав — переписан.
    changed = podborka(20, rows=[{"block": "Нектары", "sku": "1004", "name": "НЕКТАР", "cat5": 61050102}])
    ex.run(cfg(tmp_path), Portal([], [changed]), Mci([]), NOW, Sleep())
    assert not (folder / "подборка_пиво_101402.xlsx").exists()
    assert sheet(folder / "подборка_детская_101397.xlsx")[1][1][4:] == ["Нектары", "НЕКТАР", "1004"]
    assert manual.read_bytes() == b"x"
    assert events(tmp_path)[-2:] == ["podborka_written", "podborka_removed"]


def test_podborka_text_starting_with_equals_is_not_formula(tmp_path):
    ex.run(cfg(tmp_path), Portal([], [podborka(20, episode="=HYPERLINK(1)")]), Mci([]), NOW, Sleep())
    ws = sheet(tmp_path / "podborki" / "подборка_детская_101397.xlsx")[2]
    assert ws["D2"].value == "=HYPERLINK(1)" and ws["D2"].data_type == "s"


def test_podborka_bad_file_name_writes_nothing(tmp_path):
    bad = podborka(20, file="../inbox/x.xlsx")
    assert ex.run(cfg(tmp_path), Portal([], [bad]), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "inbox").iterdir()) and not list((tmp_path / "podborki").iterdir())
    assert events(tmp_path) == ["podborka_failed"]


def test_podborka_bad_file_name_does_not_stop_other_files(tmp_path):
    rows = [podborka(20, file="../inbox/x.xlsx"), podborka(20)]
    assert ex.run(cfg(tmp_path), Portal([], rows), Mci([]), NOW, Sleep()) == 1
    assert [p.name for p in (tmp_path / "podborki").iterdir()] == ["подборка_детская_101397.xlsx"]
    assert events(tmp_path) == ["podborka_failed", "podborka_written"]


def test_podborka_error_does_not_stop_lists(tmp_path):
    portal = Portal([promo()], [])
    portal.podborki_rows = RuntimeError("relation cvm.podborka_publish does not exist")
    assert ex.run(cfg(tmp_path), portal, Mci(guids(3)), NOW, Sleep()) == 1
    assert events(tmp_path) == ["podborka_failed", "exported"]


# ---------- картинки купонов (решение Елены 21.10.2026) ----------

JPEG = b"\xff\xd8\xff\xe0" + b"\x01" * 300
PNG = b"\x89PNG\r\n\x1a\n" + b"\x02" * 300


def coupon_cfg(tmp_path):
    return ex.Config(storage=tmp_path / "lists", state=tmp_path / "_state", retention_days=14,
                     stable_seconds=300, stable_tries=6, coupon_images=tmp_path / "images")


def coupon_image(tmp_path, rel, data=JPEG):
    path = tmp_path / "images" / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return hashlib.sha256(data).hexdigest()


def coupon(tmp_path, number=101416, variant=1, data=JPEG, ext="jpg", start=dt.date(2026, 10, 21), **over):
    rel = f"{number}/{number}-v{variant}-abc.{ext}"
    sha = coupon_image(tmp_path, rel, data)
    row = {"file_name": f"{number}_{start:%Y%m%d}.{ext}", "promo_number": number, "title": "Пельмени",
           "start_date": start, "end_date": start + dt.timedelta(days=6), "coupon_name": "-20% на пельмени",
           "art": {"title": "СКИДКА", "value": "20%", "sub": "на пельмени и вареники", "extra": "x"},
           "file_path": rel, "sha256": sha, "model": "gemini-test", "chosen_at": NOW}
    row.update(over)
    return row


def coupon_events(tmp_path):
    return [e for e in events(tmp_path) if e.startswith("coupon_")]


def test_coupon_written_with_passport_last_and_kept(tmp_path):
    row = coupon(tmp_path)
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep()) == 0
    folder = tmp_path / "lists" / "coupons"
    assert sorted(p.name for p in folder.iterdir()) == ["101416_20261021.jpg", "101416_20261021.json"]
    assert (folder / "101416_20261021.jpg").read_bytes() == JPEG
    assert oct(folder.stat().st_mode & 0o777) == "0o700"
    assert oct((folder / "101416_20261021.jpg").stat().st_mode & 0o777) == "0o600"
    passport = json.loads((folder / "101416_20261021.json").read_text(encoding="utf-8"))
    assert passport == {"promo": 101416, "title": "Пельмени", "start_date": "2026-10-21", "end_date": "2026-10-27",
                        "coupon_name": "-20% на пельмени",
                        "art": {"title": "СКИДКА", "value": "20%", "sub": "на пельмени и вареники"},
                        "sha256": row["sha256"], "model": "gemini-test", "file": "101416_20261021.jpg",
                        "exported_at": NOW.isoformat(timespec="seconds")}
    # паспорт записан не раньше картинки
    assert (folder / "101416_20261021.json").stat().st_mtime_ns >= (folder / "101416_20261021.jpg").stat().st_mtime_ns
    assert coupon_events(tmp_path) == ["coupon_written"]
    # Повтор без изменений — не переписывается.
    mtime = (folder / "101416_20261021.jpg").stat().st_mtime_ns
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep())
    assert coupon_events(tmp_path) == ["coupon_written"]
    assert (folder / "101416_20261021.jpg").stat().st_mtime_ns == mtime
    # Состояние — вне каталога SFTP.
    state = json.loads((tmp_path / "_state" / "coupons.json").read_text(encoding="utf-8"))
    assert state == {"101416": {"file": "101416_20261021.jpg", "sha256": row["sha256"]}}


def test_coupon_other_variant_rewrites(tmp_path):
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[coupon(tmp_path)]), Mci([]), NOW, Sleep())
    # Выбрали вариант PNG: имя меняется, старая пара уходит, новая пишется.
    row = coupon(tmp_path, variant=2, data=PNG, ext="png")
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep()) == 0
    folder = tmp_path / "lists" / "coupons"
    assert sorted(p.name for p in folder.iterdir()) == ["101416_20261021.json", "101416_20261021.png"]
    assert json.loads((folder / "101416_20261021.json").read_text(encoding="utf-8"))["sha256"] == row["sha256"]
    # Другой JPEG того же имени: файл переписывается.
    row3 = coupon(tmp_path, variant=3, data=JPEG + b"3")
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row3]), Mci([]), NOW, Sleep())
    assert sorted(p.name for p in folder.iterdir()) == ["101416_20261021.jpg", "101416_20261021.json"]
    assert (folder / "101416_20261021.jpg").read_bytes() == JPEG + b"3"
    assert coupon_events(tmp_path) == ["coupon_written"] * 3


def test_coupon_missing_file_on_sftp_rewritten(tmp_path):
    row = coupon(tmp_path)
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep())
    (tmp_path / "lists" / "coupons" / "101416_20261021.jpg").unlink()
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep())
    assert (tmp_path / "lists" / "coupons" / "101416_20261021.jpg").exists()
    assert coupon_events(tmp_path) == ["coupon_written", "coupon_written"]


def test_coupon_withdrawn_removed_foreign_files_kept(tmp_path):
    folder = tmp_path / "lists" / "coupons"
    folder.mkdir(parents=True)
    (folder / "101500_20261101.jpg").write_bytes(b"manual")      # положено руками — не наше
    (folder / "notes.txt").write_bytes(b"x")
    rows = [coupon(tmp_path), coupon(tmp_path, number=101432, start=dt.date(2026, 10, 29))]
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=rows), Mci([]), NOW, Sleep())
    assert (folder / "101432_20261029.jpg").exists()
    # 101432 больше нет в представлении (снята, перезаказана или старт больше 14 дней назад) — её пара уходит.
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=rows[:1]), Mci([]), NOW, Sleep()) == 0
    assert sorted(p.name for p in folder.iterdir()) == ["101416_20261021.jpg", "101416_20261021.json",
                                                         "101500_20261101.jpg", "notes.txt"]
    assert (folder / "101500_20261101.jpg").read_bytes() == b"manual"
    assert coupon_events(tmp_path) == ["coupon_written", "coupon_written", "coupon_removed"]


def test_coupon_path_outside_images_rejected(tmp_path):
    secret = tmp_path / "secret.jpg"
    secret.write_bytes(JPEG)
    sha = hashlib.sha256(JPEG).hexdigest()
    bad = [coupon(tmp_path, file_path="../secret.jpg", sha256=sha),
           coupon(tmp_path, number=101432, file_path=str(secret), sha256=sha)]
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=bad), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "lists" / "coupons").iterdir())
    assert coupon_events(tmp_path) == ["coupon_failed", "coupon_failed"]


def test_coupon_symlink_rejected(tmp_path):
    secret = tmp_path / "secret.jpg"
    secret.write_bytes(JPEG)
    sha = hashlib.sha256(JPEG).hexdigest()
    (tmp_path / "images" / "101416").mkdir(parents=True)
    os.symlink(secret, tmp_path / "images" / "101416" / "101416-v1-abc.jpg")       # ссылка на файл
    os.symlink(tmp_path, tmp_path / "images" / "101432")                           # ссылка на каталог
    rows = [{**coupon(tmp_path, number=101499), "promo_number": 101416, "file_name": "101416_20261021.jpg",
             "file_path": "101416/101416-v1-abc.jpg", "sha256": sha},
            {**coupon(tmp_path, number=101498), "promo_number": 101432, "file_name": "101432_20261021.jpg",
             "file_path": "101432/secret.jpg", "sha256": sha}]
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=rows), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "lists" / "coupons").iterdir())
    assert coupon_events(tmp_path) == ["coupon_failed", "coupon_failed"]


def test_coupon_sha_mismatch_failed_and_others_written(tmp_path):
    rows = [coupon(tmp_path, sha256="0" * 64), coupon(tmp_path, number=101432, start=dt.date(2026, 10, 29))]
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=rows), Mci([]), NOW, Sleep()) == 1
    assert sorted(p.name for p in (tmp_path / "lists" / "coupons").iterdir()) == [
        "101432_20261029.jpg", "101432_20261029.json"]
    assert coupon_events(tmp_path) == ["coupon_failed", "coupon_written"]
    text = (tmp_path / "_state" / "journal.jsonl").read_text(encoding="utf-8")
    assert "SHA-256 файла не совпал с базой" in text


def test_coupon_sha_mismatch_of_new_choice_removes_old(tmp_path):
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[coupon(tmp_path)]), Mci([]), NOW, Sleep())
    # Выбран другой вариант, а его файл не сходится с базой: прежняя картинка — уже не выбор, на SFTP её не оставляем.
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[coupon(tmp_path, variant=2, sha256="1" * 64)]), Mci([]), NOW,
           Sleep())
    assert not list((tmp_path / "lists" / "coupons").iterdir())
    assert coupon_events(tmp_path) == ["coupon_written", "coupon_failed", "coupon_removed"]


def test_coupon_bad_name_and_content_rejected(tmp_path):
    rows = [coupon(tmp_path, file_name="../inbox/101416_20261021.jpg"),
            coupon(tmp_path, number=101432, file_name="101433_20261021.jpg"),      # номер не тот
            coupon(tmp_path, number=101434, file_name="101434_20261022.jpg"),      # дата не старта
            coupon(tmp_path, number=101435, data=PNG)]                             # PNG под именем .jpg
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=rows), Mci([]), NOW, Sleep()) == 1
    assert not list((tmp_path / "lists" / "inbox").iterdir())
    assert not list((tmp_path / "lists" / "coupons").iterdir())
    assert coupon_events(tmp_path) == ["coupon_failed"] * 4


def test_coupon_view_missing_is_quiet_and_keeps_files(tmp_path):
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[coupon(tmp_path)]), Mci([]), NOW, Sleep())
    # Представления нет (None) — ни ошибки, ни удаления выложенного.
    assert ex.run(coupon_cfg(tmp_path), Portal([], coupons=None), Mci([]), NOW, Sleep()) == 0
    assert len(list((tmp_path / "lists" / "coupons").iterdir())) == 2
    assert coupon_events(tmp_path) == ["coupon_written"]


def test_coupon_db_error_does_not_stop_lists_or_remove(tmp_path):
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[coupon(tmp_path)]), Mci([]), NOW, Sleep())
    portal = Portal([promo()], coupons=RuntimeError("permission denied for view coupon_publish"))
    assert ex.run(coupon_cfg(tmp_path), portal, Mci(guids(3)), NOW, Sleep()) == 1
    assert len(list((tmp_path / "lists" / "coupons").iterdir())) == 2
    assert events(tmp_path)[-2:] == ["coupon_failed", "exported"]


def test_coupon_images_not_mounted_fails_without_removal(tmp_path):
    row = coupon(tmp_path)
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[row]), Mci([]), NOW, Sleep())
    no_mount = ex.Config(storage=tmp_path / "lists", state=tmp_path / "_state", retention_days=14)
    # Тот же выбор, файлы на месте — каталог картинок не нужен.
    assert ex.run(no_mount, Portal([], coupons=[row]), Mci([]), NOW, Sleep()) == 0
    # Новый выбор без каталога — ошибка, прежняя картинка снимается.
    assert ex.run(no_mount, Portal([], coupons=[coupon(tmp_path, variant=2, data=JPEG + b"2")]), Mci([]), NOW,
                  Sleep()) == 1
    assert coupon_events(tmp_path) == ["coupon_written", "coupon_failed", "coupon_removed"]


def test_coupon_stale_part_removed_foreign_part_kept(tmp_path):
    folder = tmp_path / "lists" / "coupons"
    folder.mkdir(parents=True)
    mine, foreign = folder / ".101416_20261021.jpg.part", folder / ".other.part"
    for f in (mine, foreign):
        f.write_bytes(b"x")
        os.utime(f, (NOW.timestamp() - 3600, NOW.timestamp() - 3600))
    ex.run(coupon_cfg(tmp_path), Portal([], coupons=[]), Mci([]), NOW, Sleep())
    assert not mine.exists() and foreign.exists()


# ---------- письмо о готовых списках ----------

NOTIFY = ex.Notify(to=("M.Panov@dixy.ru",), sender="cvm@mci-team.pro", host="smtp.test", port=465,
                   user="cvm@mci-team.pro", password="x")


def cfg_notify(tmp_path):
    return ex.Config(storage=tmp_path, state=tmp_path / "_state", retention_days=14,
                     stable_seconds=300, stable_tries=6, notify=NOTIFY)


class Send:
    def __init__(self, fail=False):
        self.fail, self.calls = fail, []

    def __call__(self, cfg, lists):
        self.calls.append([dict(x) for x in lists])
        if self.fail:
            raise OSError("smtp недоступен")


def test_notify_one_letter_with_all_lists_of_run(tmp_path, monkeypatch):
    send = Send()
    monkeypatch.setattr(ex, "send_notify", send)
    portal = Portal([promo(days=3, number=101419), promo(days=2, number=101420)])
    assert ex.run(cfg_notify(tmp_path), portal, Mci(guids(1234)), NOW, Sleep()) == 0
    assert len(send.calls) == 1
    lists = send.calls[0]
    assert sorted(x["promo"] for x in lists) == [101419, 101420]
    assert all(x["count"] == 1234 and x["file"] == f"{x['promo']}_{x['start_date'].replace('-', '')}.csv" for x in lists)
    text = ex.notify_text(lists)
    assert text.index("101420_") < text.index("101419_")          # по дате старта
    assert "клиентов: 1\u00a0234" in text and "Тест" in text
    assert not any(g in text for g in guids(1234))
    assert events(tmp_path)[-1] == "notified"
    # следующий прогон без новых списков — письма нет
    assert ex.run(cfg_notify(tmp_path), portal, Mci(guids(1234)), NOW + dt.timedelta(hours=1), Sleep()) == 0
    assert len(send.calls) == 1


def test_notify_failed_is_sent_next_run(tmp_path, monkeypatch):
    monkeypatch.setattr(ex, "send_notify", Send(fail=True))
    assert ex.run(cfg_notify(tmp_path), Portal([promo()]), Mci(guids(3)), NOW, Sleep()) == 1
    assert len(list((tmp_path / "inbox").glob("*.csv"))) == 1      # список выложен, несмотря на ошибку письма
    assert (tmp_path / "_state" / "notify_pending.json").exists()
    assert events(tmp_path)[-1] == "notify_failed"
    send = Send()
    monkeypatch.setattr(ex, "send_notify", send)
    assert ex.run(cfg_notify(tmp_path), Portal([promo()]), Mci(guids(3)), NOW + dt.timedelta(minutes=15), Sleep()) == 0
    assert [x["promo"] for x in send.calls[0]] == [101500]
    assert not (tmp_path / "_state" / "notify_pending.json").exists()


def test_notify_off_without_password(tmp_path, monkeypatch):
    monkeypatch.setenv("NOTIFY_TO", "M.Panov@dixy.ru")
    monkeypatch.setenv("NOTIFY_FROM", "cvm@mci-team.pro")
    monkeypatch.delenv("NOTIFY_SMTP_PASS", raising=False)
    assert ex.Notify.from_env() is None
    monkeypatch.setenv("NOTIFY_SMTP_PASS", "x")
    n = ex.Notify.from_env()
    assert n.host == "mail.hosting.reg.ru" and n.port == 465 and n.user == "cvm@mci-team.pro"
    send = Send()
    monkeypatch.setattr(ex, "send_notify", send)
    assert ex.run(cfg(tmp_path), Portal([promo()]), Mci(guids(3)), NOW, Sleep()) == 0
    assert send.calls == [] and events(tmp_path) == ["exported"]


def test_notify_portal_list_included(tmp_path, monkeypatch):
    send = Send()
    monkeypatch.setattr(ex, "send_notify", send)
    c = ex.Config(storage=tmp_path / "lists", state=tmp_path / "_state", retention_days=14,
                  portal_outbox=tmp_path / "outbox", notify=NOTIFY)
    outbox_pair(c.portal_outbox, number=500001, n=4)
    assert ex.run(c, Portal([]), Mci([]), NOW, Sleep()) == 0
    assert [(x["promo"], x["count"]) for x in send.calls[0]] == [(500001, 4)]


def test_send_notify_message(monkeypatch):
    sent = {}

    class Smtp:
        def __init__(self, host, port, timeout):
            sent["addr"] = (host, port, timeout)
        def __enter__(self):
            return self
        def __exit__(self, *a):
            return False
        def login(self, user, password):
            sent["login"] = user
        def send_message(self, msg):
            sent["msg"] = msg

    monkeypatch.setattr(ex.smtplib, "SMTP_SSL", Smtp)
    ex.send_notify(NOTIFY, [{"promo": 101419, "title": "20 бумажные влажные салфетки", "start_date": "2026-10-25",
                             "count": 2697958, "file": "101419_20261025.csv"}])
    msg = sent["msg"]
    assert sent["addr"] == ("smtp.test", 465, 20) and sent["login"] == "cvm@mci-team.pro"
    assert msg["Subject"] == "Готовы списки к загрузке в Campaign"
    assert msg["From"] == "cvm@mci-team.pro" and msg["To"] == "M.Panov@dixy.ru"
    assert "101419_20261025.csv — 101419_20 бумажные влажные салфетки, старт 25.10.2026, клиентов: 2\u00a0697\u00a0958" in msg.get_content()
