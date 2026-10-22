#!/usr/bin/env python3
"""Выгрузка списков CRM_GUID согласованных акций в хранилище на ru для загрузки в Manzana Campaign.

Процесс согласован с Еленой 08.10.2026. Запускается таймером на ru (cvm-campaign-lists.timer) в одноразовом
контейнере. Один прогон:
1. из базы портала — все акции со статусом «Согласовано» (approved, кто бы ни согласовал), которые ещё не начались
   (старт сегодня или позже по московской дате). Выгрузка сразу, как только акция согласована и список готов
   (решение Елены 08.10.2026; прежнее «за 2 дня до старта» отменено);
2. по каждой ещё не выгруженной — число клиентов скриптом Елены (sql/offer_count.sql, CONTROL_GROUP = 0);
   список готов, если число не менялось последние 5 минут (загрузка в базу не идёт): считается, через 5 минут
   ещё раз, при расхождении — снова через 5 минут, не дольше STABLE_TRIES раз, дальше — в следующем прогоне (15 минут);
   готовый список — скриптом Елены: CRM_GUID из I_PROMO_OFFER с CONTROL_GROUP = 0 (sql/offer_guids.sql);
3. результат — в inbox/<акция>_<старт>.csv, одна колонка без заголовка, рядом паспорт файла (манифест).
   Число строк файла сверяется с числом из шага 2. Каждая акция выгружается один раз.
Списки портала («Срезы базы», решения Елены 10.10.2026): портал при создании списка кладёт пару «список + паспорт» в папку
передачи (LISTS_PORTAL_OUTBOX, на ru — /srv/cvm_portal/api-data/campaign_outbox), задание проверяет её строго (имя, паспорт,
каждая строка — CRM_GUID, число строк и SHA-256) и переносит в inbox/ как обычный список. В I_PROMO_OFFER у таких акций
LOAD_TO_ML = 2 — по базе задание их не выгружает: файл уже передал портал, иначе список ушёл бы в Campaign дважды.
Файлы старше RETENTION_DAYS в sent/ и failed/ удаляются; отметка о выгрузке живёт до старта акции + RETENTION_DAYS.
Подборки товаров под согласованные пуши серий (решение Елены 20.10.2026): готовые подборки из базы портала
(cvm.podborka_publish) — файлом на акцию в podborki/, все даты акции в одном файле; файл без подборок удаляется.
Картинки купонов (решение Елены 21.10.2026: «сами купоны пока не подключили — складывай на SFTP»): выбранная во вкладке
«Купоны» портала картинка (cvm.coupon_publish) — копией в coupons/ с паспортом рядом; акция ушла из представления —
файлы удаляются.

CRM_GUID есть только в файлах списков. В журнал, манифест, состояние и вывод службы — номера акций, числа и хеши.
"""
from __future__ import annotations

import datetime as dt
import fcntl
import hashlib
import json
import os
import re
import signal
import smtplib
import sys
import stat
import time
from dataclasses import dataclass
from email.message import EmailMessage
from email.utils import formatdate, make_msgid
from pathlib import Path
from zoneinfo import ZoneInfo

MSK = ZoneInfo("Europe/Moscow")
SQL_DIR = Path(__file__).resolve().parent / "sql"


@dataclass(frozen=True)
class Config:
    storage: Path      # списки: inbox/, processing/, sent/, failed/, receipts/ — эту часть видит партнёр по SFTP
    state: Path        # журнал, отметки о выгрузке, замок — вне каталога SFTP
    retention_days: int
    stable_seconds: int = 300   # «не менялся последние 5 минут» — решение Елены 08.10.2026
    stable_tries: int = 1       # одна пауза 5 минут за прогон; не устоялось — следующая проверка через 15 минут
    deadline_seconds: int = 600   # после 10 минут прогона новые паузы не начинаются: таймер каждые 15 минут
    portal_outbox: Path | None = None   # папка передачи списков портала; None — списков портала нет
    coupon_images: Path | None = None   # каталог картинок купонов портала (только чтение); None — не смонтирован
    notify: "Notify | None" = None       # письмо о готовых списках; None — выключено

    @classmethod
    def from_env(cls) -> "Config":
        outbox = os.environ.get("LISTS_PORTAL_OUTBOX", "").strip()
        images = os.environ.get("COUPON_IMAGES_DIR", "/coupon-images").strip()
        return cls(storage=Path(os.environ.get("LISTS_STORAGE", "/data")),
                   state=Path(os.environ.get("LISTS_STATE", "/state")),
                   retention_days=int(os.environ.get("RETENTION_DAYS", "14")),
                   stable_seconds=int(os.environ.get("STABLE_SECONDS", "300")),
                   stable_tries=int(os.environ.get("STABLE_TRIES", "1")),
                   portal_outbox=Path(outbox) if outbox else None,
                   coupon_images=Path(images) if images else None,
                   notify=Notify.from_env())


@dataclass(frozen=True)
class Promo:
    number: int
    title: str
    start_date: dt.date
    approved_by: str


# ---------- хранилище ----------

class Storage:
    DIRS = ("inbox", "processing", "sent", "failed", "receipts")

    def __init__(self, root: Path, state_root: Path):
        self.root = root
        for name in self.DIRS:
            (root / name).mkdir(mode=0o700, parents=True, exist_ok=True)
        self.state_dir = state_root / "state"
        self.state_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
        self.journal_path = state_root / "journal.jsonl"
        self.exported: list[dict] = []   # выложено в inbox/ за этот прогон — для письма о готовых списках

    def path(self, name: str) -> Path:
        return self.root / name

    def journal(self, event: str, **fields) -> None:
        row = {"at": dt.datetime.now(MSK).isoformat(timespec="seconds"), "event": event, **fields}
        line = json.dumps(row, ensure_ascii=False)
        with open(self.journal_path, "a", encoding="utf-8") as f:
            f.write(line + "\n")
        print(line, flush=True)

    def exported_mark(self, base: str) -> Path:
        return self.state_dir / f"{base}.exported"


def write_list(storage: Storage, base: str, guids, expected: int, manifest: dict, folder: str = "inbox") -> dict:
    """Пишет список атомарно: .part → сверка числа строк → паспорт → переименование. Паспорт появляется последним,
    поэтому тот, кто берёт файл только вместе с паспортом, недописанного списка не увидит.
    Значения пишутся как вернул скрипт: CRM_GUID строкой, пустой CRM_GUID — пустой строкой."""
    target = storage.path(folder)
    part = target / f"{base}.csv.part"
    digest = hashlib.sha256()
    count = 0
    fd = os.open(part, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as f:
            for guid in guids:
                line = ("" if guid is None else str(guid).strip().upper()) + "\r\n"
                f.write(line)
                digest.update(line.encode("utf-8"))
                count += 1
            f.flush()
            os.fsync(f.fileno())
    except BaseException:
        # Обрыв чтения из MCI посреди списка: недописанный файл с CRM_GUID не оставляем до следующей уборки.
        part.unlink(missing_ok=True)
        raise
    if count != expected:
        part.unlink(missing_ok=True)
        raise ValueError(f"в файле {count} строк, по числу клиентов {expected}")
    manifest = {**manifest, "file": f"{base}.csv", "count": count, "sha256": digest.hexdigest()}
    mpart = target / f"{base}.manifest.json.part"
    fd = os.open(mpart, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)
        f.flush()
        os.fsync(f.fileno())
    part.replace(target / f"{base}.csv")
    mpart.replace(target / f"{base}.manifest.json")
    return manifest


def cleanup(storage: Storage, retention_days: int, now: dt.datetime) -> None:
    """Файлы старше срока удаляются: держать CRM_GUID дольше согласованных 14 дней незачем."""
    edge = now.timestamp() - retention_days * 86400
    for folder in ("sent", "failed"):
        for p in sorted(storage.path(folder).iterdir()):
            if p.is_file() and p.stat().st_mtime < edge:
                p.unlink()
                storage.journal("deleted", folder=folder, file=p.name, retention_days=retention_days)
    for p in storage.state_dir.iterdir():
        if not p.is_file():
            continue
        start = mark_start(p.name)
        # Отметку держим, пока акция может вернуться в выборку (старт сегодня или позже), а не 14 дней с выгрузки:
        # акция, согласованная за три недели до старта, иначе выгрузилась бы второй раз — её файл в sent/ к тому
        # времени уже удалён, и защита по манифесту не сработает.
        if start is not None:
            if start < now.astimezone(MSK).date() - dt.timedelta(days=retention_days):
                p.unlink()
        elif p.stat().st_mtime < edge:
            p.unlink()
    # Недописанные .part от упавшего прогона: их никто не загрузит, а GUID в них есть.
    for folder in ("inbox",):
        for p in storage.path(folder).glob("*.part"):
            if p.stat().st_mtime < now.timestamp() - 600:
                p.unlink()
                storage.journal("deleted", folder=folder, file=p.name, reason="недописанный файл старше 10 минут")


MARK_NAME = re.compile(r"\d+_(\d{8})\.exported")


def mark_start(name: str) -> dt.date | None:
    """Дата старта акции из имени отметки «<акция>_<ГГГГММДД>.exported»; чужое имя — None."""
    m = MARK_NAME.fullmatch(name)
    if not m:
        return None
    try:
        return dt.datetime.strptime(m.group(1), "%Y%m%d").date()
    except ValueError:
        return None


# ---------- списки портала ----------

PORTAL_BASE = re.compile(r"(\d{6})_(\d{8})")
PORTAL_LINE = re.compile(rb"[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\r\n")
TEXT_OK = re.compile(r"[^\x00-\x1f\x7f\u2028\u2029]{1,250}")
MANIFEST_MAX = 64 * 1024


def _open_regular(path: Path):
    """Файл из папки передачи — только обычный файл, по ссылке не открывается: в папку пишет API портала, и ссылка
    на чужой файл задания не должна попасть в обмен."""
    fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
    if not stat.S_ISREG(os.fstat(fd).st_mode):
        os.close(fd)
        raise ValueError(f"{path.name}: не обычный файл")
    return os.fdopen(fd, "rb")


def read_portal_manifest(path: Path, base: str) -> dict:
    """Паспорт списка портала: те же поля, что у паспорта задания; номер и дата старта совпадают с именем пары."""
    with _open_regular(path) as f:
        raw = f.read(MANIFEST_MAX + 1)
    if len(raw) > MANIFEST_MAX:
        raise ValueError("паспорт больше 64 КБ")
    m = json.loads(raw.decode("utf-8"))
    number, start = PORTAL_BASE.fullmatch(base).groups()
    start_iso = f"{start[:4]}-{start[4:6]}-{start[6:]}"
    dt.date.fromisoformat(start_iso)
    ok = (isinstance(m, dict) and type(m.get("promo")) is int and m["promo"] == int(number)
          and m.get("start_date") == start_iso and m.get("file") == f"{base}.csv"
          and type(m.get("count")) is int and m["count"] > 0
          and isinstance(m.get("sha256"), str) and re.fullmatch(r"[0-9a-f]{64}", m["sha256"])
          and all(isinstance(m.get(k), str) and TEXT_OK.fullmatch(m[k]) for k in ("title", "approved_by", "exported_at")))
    if not ok:
        raise ValueError("паспорт не по формату обмена")
    return {k: m[k] for k in ("promo", "title", "start_date", "approved_by", "exported_at", "count", "sha256")}


def portal_guids(f, sha256: str):
    """Строки списка портала: каждая — CRM_GUID в верхнем регистре и CRLF; в конце SHA-256 сверяется с паспортом —
    до переименования, поэтому несовпавший файл в inbox не появляется (write_list удалит .part)."""
    digest = hashlib.sha256()
    for raw in f:
        if not PORTAL_LINE.fullmatch(raw):
            raise ValueError("в списке строка не CRM_GUID формата обмена")
        digest.update(raw)
        yield raw[:-2].decode("ascii")
    if digest.hexdigest() != sha256:
        raise ValueError("SHA-256 списка не совпал с паспортом")


def take_portal_lists(cfg: Config, storage: Storage, now: dt.datetime) -> int:
    """Пары портала из папки передачи → inbox/. Уже выгруженная (отметка или паспорт в inbox/processing/sent) — пара
    убирается без второй выгрузки. Ошибка проверки — пара остаётся в папке, причина в журнале. Возвращает число ошибок."""
    outbox = cfg.portal_outbox
    if outbox is None or not outbox.is_dir():
        return 0
    errors = 0
    for mpath in sorted(outbox.glob("*.manifest.json")):
        base = mpath.name[:-len(".manifest.json")]
        csv = outbox / f"{base}.csv"
        if not PORTAL_BASE.fullmatch(base) or not csv.exists():
            continue                                   # чужое имя или список ещё пишется
        try:
            if storage.exported_mark(base).exists() or any(
                    (storage.path(f) / f"{base}.manifest.json").exists() for f in ("inbox", "processing", "sent")):
                csv.unlink()
                mpath.unlink()
                storage.journal("portal_duplicate", promo=int(base.split("_")[0]), file=f"{base}.csv")
                continue
            m = read_portal_manifest(mpath, base)
            with _open_regular(csv) as f:
                manifest = {k: m[k] for k in ("promo", "title", "start_date", "approved_by", "exported_at")}
                written = write_list(storage, base, portal_guids(f, m["sha256"]), m["count"], manifest)
        except Exception as error:  # noqa: BLE001 — одна пара не должна останавливать остальные
            storage.journal("export_failed", source="portal", file=f"{base}.csv", reason=str(error)[:300])
            errors += 1
            continue
        mark_exported(storage, base, now)
        csv.unlink()
        mpath.unlink()
        storage.exported.append({"promo": m["promo"], "title": m["title"], "start_date": m["start_date"],
                                 "count": written["count"], "file": f"{base}.csv"})
        storage.journal("exported", source="portal", promo=m["promo"], start_date=m["start_date"],
                        count=written["count"], sha256=written["sha256"], approved_by=m["approved_by"])
    return errors


# ---------- письмо о готовых списках ----------

# Решение Елены 22.10.2026: «при загрузке списков на SFTP отправляется сообщение от cvm@mci-team.pro на ящик Панова
# „Готовы списки к загрузке в Campaign“, в теле письма перечисляются списки для загрузки и кол-во клиентов в них».
# Одно письмо на прогон — все списки, выложенные в inbox/ за этот прогон (акции и списки портала). Письмо не
# отправилось — списки ждут в state/notify_pending.json и уходят следующим прогоном; на выгрузку ошибка письма не влияет.
# В письме — имена файлов, названия и числа, CRM_GUID нет. Отправка — SMTP хостинга reg.ru (SPF домена mci-team.pro
# разрешает именно его, как у писем портала в cvm_portal/auth/app/mail.py); пароль ящика — только в lists.env.
NOTIFY_SUBJECT = "Готовы списки к загрузке в Campaign"
NOTIFY_TIMEOUT = 20


@dataclass(frozen=True)
class Notify:
    to: tuple[str, ...]
    sender: str
    host: str
    port: int
    user: str
    password: str

    @classmethod
    def from_env(cls) -> "Notify | None":
        """None — письмо выключено: нет получателя, отправителя или пароля."""
        to = tuple(a.strip() for a in os.environ.get("NOTIFY_TO", "").split(",") if a.strip())
        sender = os.environ.get("NOTIFY_FROM", "").strip()
        password = os.environ.get("NOTIFY_SMTP_PASS", "")
        if not (to and sender and password):
            return None
        return cls(to=to, sender=sender, host=os.environ.get("NOTIFY_SMTP_HOST", "mail.hosting.reg.ru").strip(),
                   port=int(os.environ.get("NOTIFY_SMTP_PORT", "465")),
                   user=os.environ.get("NOTIFY_SMTP_USER", "").strip() or sender, password=password)


def notify_text(lists: list[dict]) -> str:
    lines = ["Добрый день!", "", "На SFTP (папка inbox) готовы списки к загрузке в Campaign:", ""]
    for x in sorted(lists, key=lambda x: (x["start_date"], x["promo"])):
        start = dt.date.fromisoformat(x["start_date"]).strftime("%d.%m.%Y")
        count = f"{x['count']:,}".replace(",", " ")
        lines.append(f"• {x['file']} — {x['promo']}_{x['title']}, старт {start}, клиентов: {count}")
    lines += ["", "Письмо отправлено автоматически заданием выгрузки списков."]
    return "\n".join(lines)


def send_notify(cfg: Notify, lists: list[dict]) -> None:
    msg = EmailMessage()
    msg["Subject"] = NOTIFY_SUBJECT
    msg["From"] = cfg.sender
    msg["To"] = ", ".join(cfg.to)
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=cfg.sender.rsplit("@", 1)[-1])
    msg.set_content(notify_text(lists))
    with smtplib.SMTP_SSL(cfg.host, cfg.port, timeout=NOTIFY_TIMEOUT) as smtp:
        smtp.login(cfg.user, cfg.password)
        smtp.send_message(msg)


def notify_ready(storage: Storage, cfg: Notify | None, send=None) -> int:
    """Письмо о списках, выложенных за прогон, и о тех, что не ушли письмом раньше. Письмо выключено (cfg None) —
    ничего не делается. Возвращает число ошибок."""
    if cfg is None:
        return 0
    pending_path = storage.state_dir.parent / "notify_pending.json"
    pending = json.loads(pending_path.read_text(encoding="utf-8")) if pending_path.exists() else []
    lists = pending + storage.exported
    if not lists:
        return 0
    try:
        (send or send_notify)(cfg, lists)
    except Exception as error:  # noqa: BLE001 — письмо не должно ломать выгрузку
        part = pending_path.with_name(pending_path.name + ".part")
        part.write_text(json.dumps(lists, ensure_ascii=False, indent=1), encoding="utf-8")
        part.replace(pending_path)
        storage.journal("notify_failed", promos=[x["promo"] for x in lists], reason=str(error)[:300])
        return 1
    pending_path.unlink(missing_ok=True)
    storage.journal("notified", promos=[x["promo"] for x in lists], to=list(cfg.to))
    return 0


# ---------- прогон ----------

CAMPAIGN_LEN = 50              # имя кампании в Manzana Campaign — до 50 знаков (Елена 19.10.2026)
CAMPAIGN_SUFFIX = " PUSH"
CTRL = re.compile(r"[\x00-\x1f\x7f\x85\u2028\u2029]+")


def campaign_title(promo: "Promo", mci) -> str:
    """title паспорта: принимающая сторона называет кампанию «<номер>_<title>» (ТЗ, раздел 6). Решение Елены 19.10.2026:
    имя в Manzana Campaign — до 50 знаков и не такое, как имя в Manzana Online (I_PROMO.PROMO_NAME_ML, до 40), иначе
    Manzana его не примет. title — название акции словами целиком; совпало с PROMO_NAME_ML — в конце « PUSH»."""
    prefix = f"{promo.number}_"
    fn = getattr(mci, "promo_name_ml", None)
    ml = ((fn(promo.number) if fn else None) or "").strip()
    words = CTRL.sub(" ", promo.title or "").split()
    if not words and ml.startswith(prefix):     # у акции нет названия в портале — слова из имени для Online
        words = ml[len(prefix):].split()
    room = CAMPAIGN_LEN - len(prefix)
    out: list[str] = []
    for w in words:
        if len(" ".join(out + [w])) > room:
            break
        out.append(w)
    if not out and words:
        out = [words[0][:room]]
    title = " ".join(out)
    if ml and (prefix + title).casefold() == ml.casefold():
        title = title[:room - len(CAMPAIGN_SUFFIX)].rstrip() + CAMPAIGN_SUFFIX
    return title


def mark_exported(storage: Storage, base: str, now: dt.datetime) -> None:
    """Отметка «акция выгружена» — атомарно: .tmp → fsync → replace."""
    mark = storage.exported_mark(base)
    tmp = mark.with_name(mark.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(now.isoformat(timespec="seconds") + "\n")
        f.flush()
        os.fsync(f.fileno())
    tmp.replace(mark)


def stable_counts(cfg: Config, mci, promos: list[int], sleep) -> tuple[dict[int, int | None], dict[int, str], set[int]]:
    """Число клиентов каждой акции, если оно не менялось stable_seconds (загрузка списка в базу не идёт); None —
    список пуст или ещё загружается. Акции ждут вместе: одна пауза на всех, а не по 5 минут на каждую, — иначе
    при десятке акций прогон не укладывается в 15 минут. Второй словарь — ошибки подсчёта по акциям; множество —
    списки портала (строки с LOAD_TO_ML = 2): их не ждём и не выгружаем."""
    result: dict[int, int | None] = {n: None for n in promos}
    errors: dict[int, str] = {}
    portal: set[int] = set()

    def count(n: int) -> int | None:
        try:
            rows, portal_rows = mci.offer_count(n)
        except Exception as error:  # noqa: BLE001 — ошибка одной акции не останавливает остальные
            errors[n] = str(error)
            return None
        if portal_rows:
            portal.add(n)
        return rows

    # Часы — с начала подсчёта: при блокировке ночной заливки первые подсчёты сами идут минутами.
    started = time.monotonic()
    before = {n: count(n) for n in promos}
    pending = [n for n in promos if before[n] and n not in portal]  # 0 — списка нет; None — ошибка
    for _ in range(cfg.stable_tries):
        if not pending or time.monotonic() - started + cfg.stable_seconds > cfg.deadline_seconds:
            break
        sleep(cfg.stable_seconds)
        still = []
        for n in pending:
            after = count(n)
            if after is None or after == 0:
                continue
            if after == before[n]:
                result[n] = after
            else:
                before[n] = after
                still.append(n)
        pending = still
    return result, errors, portal


# ---------- подборки ----------

# Подборки товаров под согласованные пуши серий (решение Елены 20.10.2026): «после согласования коммуникации обязательно
# создаётся файл с подборкой товаров под акцию и выкладывается на sftp; при согласовании каждой последующей коммуникации
# добавляются её подборки в соответствующий файл». Состав собирает ферма (cvm_platform/tools/comm-writer/podborka.py),
# здесь — только файл: один на акцию, все даты в нём. Клиентов в подборке нет — номера акций, названия товаров и групп.
PODBORKA_DIR = "podborki"
PODBORKA_NAME = re.compile(r"подборка_[0-9a-zа-яё_]{1,40}_\d{6}\.xlsx")
PODBORKA_HEAD = ("ID акции", "Название акции", "Дата", "Эпизод", "Блок-категория", "SKU название", "SKU ID")
MONTHS = ("январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь",
          "декабрь")


def podborka_files(rows: list[dict]) -> tuple[dict[str, list[tuple]], list[str]]:
    """Строки представления cvm.podborka_publish → ({файл: строки листа} по датам, имена не по образцу). Чужое имя
    файла не берётся, но и остальные подборки из-за него не встают (ревью гейта 20.10.2026)."""
    files: dict[str, list[tuple]] = {}
    bad: list[str] = []
    for r in sorted(rows, key=lambda r: (r["file_name"], r["send_date"])):
        if not PODBORKA_NAME.fullmatch(r["file_name"]):
            bad.append(str(r["file_name"]))
            continue
        day = f"{r['send_date']:%d.%m.%Y}"
        for item in r["rows"] or []:
            files.setdefault(r["file_name"], []).append(
                (int(r["promo_number"]), r["series_title"], day, r["episode"] or "", item["block"], item["name"],
                 str(item["sku"])))
    return files, bad


def write_podborka(path: Path, rows: list[tuple]) -> None:
    """Файл по образцу выложенных: лист «SKU по акциям <месяц>», заголовок жирным. Строки пишутся текстом, даже если
    начинаются с «=», — формулой в файле партнёра ничего не станет. Запись атомарно: .part → переименование."""
    import io
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"SKU по акциям {MONTHS[int(rows[0][2][3:5]) - 1]}"
    ws.append(PODBORKA_HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append(row)
        for c in ws[ws.max_row]:
            if isinstance(c.value, str):
                c.data_type = "s"
    for col in "ABCDEFG":
        ws.column_dimensions[col].width = 13
    buf = io.BytesIO()
    wb.save(buf)
    # Скрытое имя — забирающая команда не примет недописанный файл за готовый; O_NOFOLLOW — ссылку, подложенную
    # в папку по SFTP, не открываем (ревью гейта 20.10.2026).
    part = path.with_name(f".{path.name}.part")
    fd = os.open(part, os.O_WRONLY | os.O_CREAT | os.O_TRUNC | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, "wb") as f:
        f.write(buf.getvalue())
        f.flush()
        os.fsync(f.fileno())
    part.replace(path)


def publish_podborki(storage: Storage, portal) -> int:
    """Файлы подборок на SFTP (lists/podborki/). Файл переписывается, только когда изменился его состав; файл, у которого
    не осталось ни одной готовой подборки (пуш больше не согласован), удаляется. Задание трогает только свои файлы —
    те, что когда-то записало само (state/podborki.json): разложенное руками, например «подборки_октябрь_7.xlsx», не его."""
    folder = storage.path(PODBORKA_DIR)
    folder.mkdir(mode=0o700, exist_ok=True)
    state_path = storage.state_dir.parent / "podborki.json"
    known: dict[str, str] = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    files, bad = podborka_files(portal.podborki())
    for name in bad:
        storage.journal("podborka_failed", file=name[:120], reason="имя файла подборки не по образцу")
    errors = len(bad)
    for name, rows in files.items():
        digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode("utf-8")).hexdigest()
        if known.get(name) == digest and (folder / name).exists():
            continue
        try:
            write_podborka(folder / name, rows)
        except Exception as error:  # noqa: BLE001 — один файл не должен останавливать остальные
            storage.journal("podborka_failed", file=name, reason=str(error)[:300])
            errors += 1
            continue
        known[name] = digest
        storage.journal("podborka_written", file=name, rows=len(rows), dates=len({r[2] for r in rows}))
    for name in sorted(set(known) - set(files)):
        (folder / name).unlink(missing_ok=True)
        del known[name]
        storage.journal("podborka_removed", file=name)
    part = state_path.with_name(state_path.name + ".part")
    part.write_text(json.dumps(known, ensure_ascii=False, indent=1), encoding="utf-8")
    part.replace(state_path)
    return errors


# ---------- картинки купонов ----------

# Решение Елены 21.10.2026: «сами купоны [в Manzana] пока не подключили — складывай на SFTP». Картинку рисует ферма
# через шлюз портала (db/gen_channel.py, файлы — /srv/cvm_portal/coupon_images на ru), человек выбирает вариант во
# вкладке «Купоны». Здесь — только копия выбранного файла и паспорт рядом. Что выкладывается (согласованная акция,
# выбранный вариант, старт не раньше 14 дней назад) — решает представление cvm.coupon_publish (миграция портала
# db/031_coupon_images.sql), задание его зеркалит. Клиентов в картинках и паспортах нет.
COUPON_DIR = "coupons"
COUPON_NAME = re.compile(r"(\d{6,9})_(\d{8})\.(jpg|png|webp)")
COUPON_SHA = re.compile(r"[0-9a-f]{64}")
COUPON_PART = re.compile(r"\.\d{6,9}_\d{8}\.(jpg|png|webp|json)\.part")
COUPON_MAX_BYTES = 5 * 1024 * 1024   # шлюз принимает до 700 КБ; больше — не картинка шлюза
COUPON_MAGIC = {"jpg": (b"\xff\xd8\xff",), "png": (b"\x89PNG\r\n\x1a\n",), "webp": (b"RIFF",)}
COUPON_ART_KEYS = ("title", "value", "sub")


def coupon_base(name: str) -> str:
    return name.rsplit(".", 1)[0]


def read_coupon_image(root: Path | None, file_path: str, sha256: str, ext: str) -> bytes:
    """Файл выбранной картинки из каталога картинок портала. Путь — только относительный внутри каталога: без «..»,
    без абсолютного пути и без ссылок ни в одном звене (resolve должен дать тот же путь), последнее звено
    открывается с O_NOFOLLOW; только обычный файл не больше COUPON_MAX_BYTES; SHA-256 содержимого — как в базе,
    сигнатура — по расширению. Иначе ValueError — картинка не выкладывается."""
    if root is None or not root.is_dir():
        raise ValueError("каталог картинок купонов не смонтирован")
    root = root.resolve()
    rel = Path(file_path or "")
    if not file_path or rel.is_absolute() or ".." in rel.parts:
        raise ValueError("путь картинки вне каталога картинок")
    path = root / rel
    if path.resolve() != path or not path.resolve().is_relative_to(root):
        raise ValueError("путь картинки вне каталога картинок или через ссылку")
    if path.suffix.lower().lstrip(".") != ext:
        raise ValueError("расширение файла не совпадает с именем на SFTP")
    fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
    with os.fdopen(fd, "rb") as f:
        if not stat.S_ISREG(os.fstat(f.fileno()).st_mode):
            raise ValueError("картинка — не обычный файл")
        data = f.read(COUPON_MAX_BYTES + 1)
    if len(data) > COUPON_MAX_BYTES:
        raise ValueError("картинка больше 5 МБ")
    if not data.startswith(COUPON_MAGIC[ext]) or (ext == "webp" and data[8:12] != b"WEBP"):
        raise ValueError("содержимое файла не совпадает с расширением")
    if hashlib.sha256(data).hexdigest() != sha256:
        raise ValueError("SHA-256 файла не совпал с базой")
    return data


def write_atomic(path: Path, data: bytes) -> None:
    """Скрытое имя .<имя>.part — забирающая сторона не примет недописанный файл за готовый; O_NOFOLLOW — ссылку,
    подложенную в папку по SFTP, не открываем; потом fsync и переименование."""
    part = path.with_name(f".{path.name}.part")
    fd = os.open(part, os.O_WRONLY | os.O_CREAT | os.O_TRUNC | os.O_NOFOLLOW, 0o600)
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        part.replace(path)
    except BaseException:
        part.unlink(missing_ok=True)
        raise


def coupon_passport(row: dict, name: str, now: dt.datetime) -> bytes:
    """Паспорт картинки: акция, даты, название купона, текст картинки, контроль файла. Без клиентов."""
    art = row.get("art")
    art = {k: str(art.get(k) or "") for k in COUPON_ART_KEYS} if isinstance(art, dict) else None
    passport = {"promo": int(row["promo_number"]), "title": row.get("title") or "",
                "start_date": str(row["start_date"]), "end_date": str(row["end_date"]) if row.get("end_date") else None,
                "coupon_name": row.get("coupon_name") or "", "art": art, "sha256": row["sha256"],
                "model": row.get("model") or "", "file": name, "exported_at": now.isoformat(timespec="seconds")}
    return json.dumps(passport, ensure_ascii=False, indent=1).encode("utf-8")


def remove_coupon(folder: Path, entry: dict) -> None:
    """Сначала паспорт, потом картинку: без паспорта забирающая сторона файл уже не берёт."""
    (folder / f"{coupon_base(entry['file'])}.json").unlink(missing_ok=True)
    (folder / entry["file"]).unlink(missing_ok=True)


def publish_coupons(storage: Storage, portal, images: Path | None, now: dt.datetime) -> int:
    """Картинки купонов на SFTP (lists/coupons/): <номер>_<ГГГГММДД>.<jpg|png|webp> и паспорт <номер>_<ГГГГММДД>.json,
    паспорт — последним. Переписывается, только когда сменился sha256 (выбран другой вариант) или файл пропал; акции нет
    в представлении (снята, купон перезаказан, старт больше 14 дней назад) — её файлы удаляются. Задание трогает только
    свои файлы — записанные им самим (state coupons.json). Представления нет (миграция портала не применена) — ничего не
    делается и ничего не удаляется. Возвращает число ошибок."""
    rows = portal.coupons()
    if rows is None:
        return 0
    folder = storage.path(COUPON_DIR)
    folder.mkdir(mode=0o700, exist_ok=True)
    state_path = storage.state_dir.parent / "coupons.json"
    known: dict[str, dict] = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    errors = 0
    seen: set[str] = set()
    for row in rows:
        name = str(row.get("file_name") or "")
        number = row.get("promo_number")
        key = str(number)
        seen.add(key)
        old = known.get(key)
        try:
            m = COUPON_NAME.fullmatch(name)
            start = row.get("start_date")
            if not m or m.group(1) != key or not isinstance(start, dt.date) or m.group(2) != f"{start:%Y%m%d}":
                raise ValueError("имя файла картинки не по образцу")
            sha = str(row.get("sha256") or "").lower()
            if not COUPON_SHA.fullmatch(sha):
                raise ValueError("нет SHA-256 картинки в базе")
            passport = folder / f"{coupon_base(name)}.json"
            if (old and old.get("file") == name and old.get("sha256") == sha
                    and (folder / name).exists() and passport.exists()):
                continue
            data = read_coupon_image(images, str(row.get("file_path") or ""), sha, m.group(3))
            # Старый паспорт — первым: пока пишется новая картинка, пара «новая картинка + старый паспорт» не видна.
            if old and old.get("file") != name:
                remove_coupon(folder, old)
            elif old:
                passport.unlink(missing_ok=True)
            write_atomic(folder / name, data)
            write_atomic(passport, coupon_passport({**row, "sha256": sha}, name, now))
        except Exception as error:  # noqa: BLE001 — одна картинка не должна останавливать остальные
            storage.journal("coupon_failed", promo=number if isinstance(number, int) else None, file=name[:120],
                            reason=str(error)[:300])
            errors += 1
            # Старая картинка — прежний выбор, сейчас выбран другой вариант: показывать её дальше хуже, чем ничего.
            if old and old.get("sha256") != str(row.get("sha256") or "").lower():
                remove_coupon(folder, old)
                del known[key]
                storage.journal("coupon_removed", promo=number, file=old["file"], reason="выбран другой вариант")
            continue
        known[key] = {"file": name, "sha256": sha}
        storage.journal("coupon_written", promo=number, file=name, sha256=sha)
    for key in sorted(set(known) - seen):
        remove_coupon(folder, known[key])
        storage.journal("coupon_removed", promo=int(key), file=known[key]["file"])
        del known[key]
    # Недописанные .part своих имён от упавшего прогона.
    for p in folder.glob(".*.part"):
        if COUPON_PART.fullmatch(p.name) and p.lstat().st_mtime < now.timestamp() - 600:
            p.unlink()
    part = state_path.with_name(state_path.name + ".part")
    part.write_text(json.dumps(known, ensure_ascii=False, indent=1), encoding="utf-8")
    part.replace(state_path)
    return errors


def run(cfg: Config, portal, mci, now: dt.datetime | None = None, sleep=time.sleep) -> int:
    """portal и mci — объекты с методами из PortalDb и MciDb (в тестах — подставные). Возвращает код выхода."""
    now = now or dt.datetime.now(MSK)
    today = now.astimezone(MSK).date()
    storage = Storage(cfg.storage, cfg.state)
    errors = 0
    # Уборка и в начале: прогон, оборванный по таймауту, до уборки в конце не доходит, а .part с GUID лежит в inbox/.
    cleanup(storage, cfg.retention_days, now)
    # Списки портала — первыми: без подсчётов и пауз, портал передал их сразу при создании.
    errors += take_portal_lists(cfg, storage, now)
    # Подборки — до списков акций: без пауз, и падение базы портала на них не мешает спискам (откат в PortalDb.podborki).
    try:
        errors += publish_podborki(storage, portal)
    except Exception as error:  # noqa: BLE001
        storage.journal("podborka_failed", reason=str(error)[:300])
        errors += 1
    # Картинки купонов — сразу после подборок, тоже без пауз; ошибка базы откатывается в PortalDb.coupons.
    try:
        errors += publish_coupons(storage, portal, cfg.coupon_images, now)
    except Exception as error:  # noqa: BLE001
        storage.journal("coupon_failed", reason=str(error)[:300])
        errors += 1

    todo = []
    for promo in portal.approved_promos(today):
        base = f"{promo.number}_{promo.start_date:%Y%m%d}"
        if storage.exported_mark(base).exists():
            continue
        if any((storage.path(f) / f"{base}.manifest.json").exists() for f in ("inbox", "processing", "sent")):
            # Файл выгружен, а отметку прогон записать не успел (оборван) — второй раз не выгружаем, иначе дубль в Campaign.
            mark_exported(storage, base, now)
            continue
        todo.append((promo, base))
    counts, count_errors, portal_lists = stable_counts(cfg, mci, sorted({p.number for p, _ in todo}), sleep)

    for promo, base in todo:
        if promo.number in count_errors:
            storage.journal("export_failed", promo=promo.number, start_date=str(promo.start_date),
                            reason=count_errors[promo.number])
            errors += 1
            continue
        if promo.number in portal_lists:
            # LOAD_TO_ML = 2: список портала, файл в обмен передал сам портал — второй раз не выгружаем (решение Елены 10.10.2026)
            mark_exported(storage, base, now)
            storage.journal("skipped_portal_list", promo=promo.number, start_date=str(promo.start_date))
            continue
        count = counts[promo.number]
        if count is None:
            # Списка ещё нет или он загружается — следующий прогон через 15 минут.
            storage.journal("not_ready", promo=promo.number, start_date=str(promo.start_date))
            continue
        try:
            manifest = {"promo": promo.number, "title": campaign_title(promo, mci), "start_date": str(promo.start_date),
                        "approved_by": promo.approved_by, "exported_at": now.isoformat(timespec="seconds")}
            written = write_list(storage, base, mci.offer_guids(promo.number), count, manifest)
        except Exception as error:  # noqa: BLE001 — одна акция не должна останавливать остальные
            storage.journal("export_failed", promo=promo.number, start_date=str(promo.start_date), reason=str(error))
            errors += 1
            continue
        mark_exported(storage, base, now)
        storage.exported.append({"promo": promo.number, "title": manifest["title"], "start_date": str(promo.start_date),
                                 "count": written["count"], "file": f"{base}.csv"})
        storage.journal("exported", promo=promo.number, start_date=str(promo.start_date), count=written["count"],
                        sha256=written["sha256"], approved_by=promo.approved_by)

    errors += notify_ready(storage, cfg.notify)
    cleanup(storage, cfg.retention_days, now)
    return 1 if errors else 0


# ---------- источники ----------

class PortalDb:
    """База портала: только чтение cvm.promo ролью cvm_lists_ro (грант на нужные колонки)."""

    def __init__(self):
        import psycopg
        self.conn = psycopg.connect(host=os.environ.get("PG_HOST", "cvm-db"), port=os.environ.get("PG_PORT", "5432"),
                                    dbname=os.environ["PG_DB"], user=os.environ["PG_USER"],
                                    password=os.environ["PG_PASSWORD"], connect_timeout=15)
        self.conn.read_only = True

    def approved_promos(self, start_from: dt.date) -> list[Promo]:
        sql = (SQL_DIR / "approved_promos.sql").read_text(encoding="utf-8")
        with self.conn.cursor() as cur:
            cur.execute(sql, {"start_from": start_from})
            return [Promo(number=r[0], title=r[1], start_date=r[2], approved_by=r[3]) for r in cur.fetchall()]

    def podborki(self) -> list[dict]:
        """Готовые подборки (представление cvm.podborka_publish, миграция портала db/025): только поля файла."""
        sql = ("select file_name, promo_number, series_title, send_date, episode, rows "
               "from cvm.podborka_publish order by file_name, send_date")
        try:
            with self.conn.cursor() as cur:
                cur.execute(sql)
                names = [d.name for d in cur.description]
                return [dict(zip(names, r)) for r in cur.fetchall()]
        except Exception:
            # Ошибка оставила бы транзакцию прерванной, и согласованные акции ниже уже не прочитались бы.
            self.conn.rollback()
            raise

    def coupons(self) -> list[dict] | None:
        """Выбранные картинки купонов (представление cvm.coupon_publish, миграция портала db/031). Представления нет —
        миграция ещё не применена: None, это не ошибка задания и не повод удалять выложенное."""
        import psycopg
        sql = ("select file_name, promo_number, title, start_date, end_date, coupon_name, art, file_path, sha256, model, "
               "chosen_at from cvm.coupon_publish order by promo_number")
        try:
            with self.conn.cursor() as cur:
                cur.execute(sql)
                names = [d.name for d in cur.description]
                return [dict(zip(names, r)) for r in cur.fetchall()]
        except psycopg.errors.UndefinedTable:
            self.conn.rollback()
            return None
        except Exception:
            self.conn.rollback()
            raise


class MciDb:
    """mci_model: учётка cvm_api (только чтение) из api.env — решение Елены 08.10.2026."""

    def __init__(self):
        import pyodbc
        # Как в API портала (api/app/config.py, _bool): пусто — по умолчанию да.
        raw = os.environ.get("MSSQL_TRUST_CERT", "").strip().lower()
        trust = "yes" if not raw or raw in ("1", "true", "yes", "y", "on") else "no"
        cs = ("DRIVER={ODBC Driver 18 for SQL Server};"
              f"SERVER={os.environ['MSSQL_HOST']},{os.environ.get('MSSQL_PORT', '1433')};"
              f"DATABASE={os.environ['MSSQL_DB']};UID={os.environ['MSSQL_USER']};PWD={os.environ['MSSQL_PASSWORD']};"
              f"Encrypt={os.environ.get('MSSQL_ENCRYPT', 'yes')};TrustServerCertificate={trust};"
              "ApplicationIntent=ReadOnly;APP=cvm-campaign-lists")
        self.conn = pyodbc.connect(cs, timeout=20, autocommit=True)
        self.conn.timeout = 300  # таймаут запроса: с отбором по ID_PROMO оба запроса на колоночном индексе идут секунды
        # Блокировку ночной заливки ждём не дольше минуты; дальше ошибка по акции и повтор в следующем прогоне.
        self.conn.execute("SET LOCK_TIMEOUT 60000")

    def _sql(self, name: str) -> str:
        return (SQL_DIR / name).read_text(encoding="utf-8")

    def offer_count(self, promo: int) -> tuple[int, int]:
        """(клиентов без КГ, из них строк списка портала с LOAD_TO_ML = 2)."""
        cur = self.conn.cursor()
        try:
            cur.execute(self._sql("offer_count.sql"), promo)
            rows, portal_rows = cur.fetchone()
            return int(rows), int(portal_rows)
        finally:
            cur.close()

    def offer_guids(self, promo: int):
        return self._stream("offer_guids.sql", promo)

    def promo_name_ml(self, promo: int) -> str | None:
        """I_PROMO.PROMO_NAME_ML — имя списка для Manzana Online; нет строки или колонки пусто — None."""
        cur = self.conn.cursor()
        try:
            cur.execute(self._sql("promo_name_ml.sql"), promo)
            row = cur.fetchone()
            return row[0] if row and row[0] else None
        finally:
            cur.close()

    def _stream(self, name: str, *params):
        cur = self.conn.cursor()
        try:
            cur.execute(self._sql(name), *params)
            while True:
                rows = cur.fetchmany(50000)
                if not rows:
                    break
                for (guid,) in rows:
                    yield guid
        finally:
            # Без MARS незакрытый курсор после ошибки даёт «Connection is busy» на следующей акции.
            cur.close()


def main() -> int:
    # SIGTERM от docker stop / systemd: выйти штатно, чтобы сработали finally и удаление .part.
    signal.signal(signal.SIGTERM, lambda *_: sys.exit(143))
    cfg = Config.from_env()
    os.umask(0o077)
    cfg.storage.mkdir(parents=True, exist_ok=True)
    cfg.state.mkdir(parents=True, exist_ok=True)
    with open(cfg.state / ".lock", "w") as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            print("прогон уже идёт, этот пропущен", flush=True)
            return 0
        return run(cfg, PortalDb(), MciDb())


if __name__ == "__main__":
    sys.exit(main())
