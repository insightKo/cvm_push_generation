#!/usr/bin/env python3
"""
Сборка полного справочника ассортимента ДИКСИ для портала CVM.

Источники (только чтение, ничего не правим):
  data/ассортимент.xlsx                      — SKU × город + метрики (лист «Лист1»)
  data/I_PRODUCT_SCORE_NEW.xlsx              — внутренний ID_PRODUCT и имена cat5 (лист «выгрузка»)
  ../dixy-dashboard/данные/I_MISSION.csv     — имена cat5 (добор недостающих)
  ../dixy-dashboard/данные/категории.xlsx    — имена cat5 (лист «Лист3»)
  ../dixy-dashboard/sql/справочник_ext_названия.sql — ручные имена cat5 (блок MANUAL)
  ../dixy-dashboard/coffee_data.js           — имена cat4 из I_PRODUCT (через внутренний ID_CATEGORY_5)
  ../dixy-dashboard/данные/Тренды по новым.xlsx — связка «имя cat4 → имя cat3» (лист «Лист3»)

ВАЖНО (решение Елены 09.09.2026): это сборка «на пока», из файловых выгрузок.
Дальше ассортимент берём из МСИ — таблицы I_PRODUCT и I_ASSORTMANF, оттуда же родные
названия категорий всех уровней; файловую сборку тогда выключаем.

Результат (output/assortment/):
  ассортимент_полный.xlsx  — 3 листа: Категории, Товары, Товары по городам
  category.csv, product.csv, product_city.csv — под \copy в Postgres портала
  (DDL и загрузка на ru — рядом: 004_assortment.sql и load_to_ru.sh)

Запуск: python3 output/build_assortment.py
"""
import collections
import csv
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DASH = ROOT.parent / 'dixy-dashboard'
OUT = ROOT / 'output' / 'assortment'
OUT.mkdir(parents=True, exist_ok=True)

SRC_ASSORT = ROOT / 'data' / 'ассортимент.xlsx'
SRC_SCORE = ROOT / 'data' / 'I_PRODUCT_SCORE_NEW.xlsx'
SRC_MISSION = DASH / 'данные' / 'I_MISSION.csv'
SRC_CATS = DASH / 'данные' / 'категории.xlsx'
SRC_MANUAL = DASH / 'sql' / 'справочник_ext_названия.sql'
SRC_COFFEE = DASH / 'coffee_data.js'
SRC_TRENDS = DASH / 'данные' / 'Тренды по новым.xlsx'

BAD_NAMES = {'', 'N/A', 'N-A', 'NULL', 'МЕППИНГ'}

METRICS = ['COUNT_SHOPS', 'COUNT_CHECK', 'REAL_SALE', 'PURCHASE_TIME', 'AVG_PRICE',
           'PENETRATION_CHECK', 'PENETRATION_BUDGET', 'COUNT_CHECK_WEEKEND', 'COUNT_CHECK_WORKDAY',
           'COUNT_CHECK_FESTADAY', 'COUNT_CHECK_MON', 'COUNT_CHECK_TUE', 'COUNT_CHECK_WED',
           'COUNT_CHECK_THU', 'COUNT_CHECK_FRI', 'COUNT_CHECK_ST', 'COUNT_CHECK_SU',
           'PAIRED_BUDGET', 'PAIRED_SKU']


def num(v):
    """Число как есть; строка вида «01,10,5382» — восстановленное Excel-ом в дату число 1.105382."""
    if v is None or v == '':
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None
    s = str(v).strip()
    parts = s.split(',')
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        return float(parts[0] + '.' + parts[1] + parts[2]), s
    try:
        return float(s.replace(',', '.')), s
    except ValueError:
        return None, s


# ── 1. Ассортимент: SKU × город ────────────────────────────────────────────
wb = openpyxl.load_workbook(SRC_ASSORT, read_only=True)
rows = wb['Лист1'].iter_rows(values_only=True)
header = [c for c in next(rows) if c]
idx = {name: header.index(name) for name in header}

city_rows = []
products = {}
cat5_parent, cat4_parent = {}, {}
for r in rows:
    if r[0] is None:
        continue
    sku = str(r[idx['ID_PRODUCT_EXTERNAL']]).strip()
    c5, c4, c3 = (str(r[idx[k]]).strip() for k in ('ID_CATEGORY_5_ext', 'ID_CATEGORY_4_ext', 'ID_CATEGORY_3_ext'))
    cat5_parent[c5], cat4_parent[c4] = c4, c3
    products.setdefault(sku, {
        'sku': sku,
        'sku_code': sku[:-2] if sku.endswith('_0') else sku,
        'product_name': str(r[idx['PRODUCT_NAME']]).strip(),
        'cat5_code': c5, 'cat4_code': c4, 'cat3_code': c3,
    })
    row = {'sku': sku, 'city': str(r[idx['CITY_REPORT']]).strip()}
    for m in METRICS:
        val, raw = num(r[idx[m]])
        row[m.lower()] = val
        if m == 'PAIRED_SKU':
            row['paired_sku_raw'] = raw
    city_rows.append(row)

# ── 2. Имена cat5: четыре источника, приоритет от самого свежего ───────────
def score_names():
    w = openpyxl.load_workbook(SRC_SCORE, read_only=True)
    it = w['выгрузка'].iter_rows(values_only=True)
    next(it)
    names, int2ext, ext2id = {}, {}, {}
    for r in it:
        if r[3] is None:
            continue
        ext = str(r[3]).strip()
        if r[4]:
            names.setdefault(ext, str(r[4]).strip())
        if r[2] is not None:
            int2ext.setdefault(int(r[2]), ext)
        if r[1] is not None and r[0] is not None:
            ext2id.setdefault(str(r[1]).strip(), int(r[0]))
    return names, int2ext, ext2id


def mission_names():
    names = {}
    with open(SRC_MISSION, encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            code, name = (r.get('ID_CATEGORY_5_ext') or '').strip(), (r.get('I_CATEGORY_5') or '').strip()
            if code and name:
                names.setdefault(code, name)
    return names


def dashboard_names():
    w = openpyxl.load_workbook(SRC_CATS, read_only=True)
    it = w['Лист3'].iter_rows(values_only=True)
    next(it)
    names = {}
    for r in it:
        if r[0] is not None and r[1]:
            names.setdefault(str(r[0]).strip(), str(r[1]).strip())
    return names


def manual_names():
    txt = SRC_MANUAL.read_text(encoding='utf-8')
    return {c: n for c, n in re.findall(r"\(N'(\d+)',\s*N'([^']*)'", txt) if n}


NAME5, INT2EXT, EXT2ID = score_names()
SRC_ORDER = [
    (NAME5, 'МСИ I_PRODUCT_SCORE'),
    (mission_names(), 'МСИ I_MISSION'),
    (dashboard_names(), 'справочник дашборда'),
    (manual_names(), 'ручной справочник (MANUAL)'),
]

# ── 3. Имена cat5 (по кодам), дальше от них пляшут cat4 и cat3 ─────────────
sku_per_cat5 = collections.Counter(p['cat5_code'] for p in products.values())
name5, src5, conflict5 = {}, {}, {}
for code in cat5_parent:
    found = [(t[code].strip(), label) for t, label in SRC_ORDER
             if t.get(code) and t[code].strip().upper() not in BAD_NAMES]
    if found:
        name5[code], src5[code] = found[0]
        conflict5[code] = len({n for n, _ in found}) > 1

# ── 4. Имена cat4: I_CATEGORY_4 из I_PRODUCT ───────────────────────────────
# Два пути к одному и тому же полю: по внутреннему ID_CATEGORY_5 (точный) и по имени cat5
# (добирает коды, которых нет в мостике внутренних id). Голос весит числом SKU под cat5.
coffee = SRC_COFFEE.read_text(encoding='utf-8')
by_id = re.findall(r'"ID_CATEGORY_5":\s*(\d+),\s*"I_CATEGORY_5":\s*"[^"]*",\s*"I_CATEGORY_4":\s*"([^"]*)"', coffee)
by_name = re.findall(r'"I_CATEGORY_5":\s*"([^"]*)",\s*"I_CATEGORY_4":\s*"([^"]*)"', coffee)

cat4_votes = collections.defaultdict(collections.Counter)
for cid, n4 in by_id:
    ext5 = INT2EXT.get(int(cid))
    c4 = cat5_parent.get(ext5)
    if c4 and n4.strip().upper() not in BAD_NAMES:
        cat4_votes[c4][n4.strip()] += sku_per_cat5.get(ext5, 1)

n5_to_n4 = {}
for n5, n4 in by_name:
    if n5.strip() and n4.strip().upper() not in BAD_NAMES:
        n5_to_n4.setdefault(n5.strip(), n4.strip())
for code, n5 in name5.items():
    n4 = n5_to_n4.get(n5)
    if n4:
        cat4_votes[cat5_parent[code]][n4] += sku_per_cat5.get(code, 1)

CAT4_NAME = {c4: v.most_common(1)[0][0] for c4, v in cat4_votes.items()}

# ── 5. Имена cat3: связка «имя cat4 → имя cat3» из трендов дашборда ────────
# Кода cat3 в МСИ нет ни в одной таблице — есть только имя рядом с именем cat4.
# Поэтому имя коду cat3 назначаем голосованием его же детей-cat4, вес — число SKU.
wt = openpyxl.load_workbook(SRC_TRENDS, read_only=True)
it = wt['Лист3'].iter_rows(values_only=True)
next(it)
n4_to_n3 = {}
for r in it:
    if r[4] and r[5] and str(r[4]).strip().upper() not in BAD_NAMES:
        n4_to_n3.setdefault(str(r[5]).strip(), str(r[4]).strip())

sku_per_cat4 = collections.Counter(p['cat4_code'] for p in products.values())
cat3_votes = collections.defaultdict(collections.Counter)
for c4, n4 in CAT4_NAME.items():
    n3 = n4_to_n3.get(n4)
    if n3:
        cat3_votes[cat4_parent[c4]][n3] += sku_per_cat4.get(c4, 1)

# ── 6. Справочник категорий: три уровня одной таблицей ─────────────────────
sku_cnt = collections.Counter()
for p in products.values():
    sku_cnt[p['cat5_code']] += 1
    sku_cnt[p['cat4_code']] += 1
    sku_cnt[p['cat3_code']] += 1

categories = []
for c3 in sorted(set(cat4_parent.values())):
    votes = cat3_votes.get(c3)
    name = votes.most_common(1)[0][0] if votes else None
    categories.append({'code': c3, 'level': 3, 'parent_code': None, 'name': name,
                       'name_source': 'МСИ I_CATEGORY_3 (через имена cat4)' if name else None,
                       'has_name_conflict': bool(votes and len(votes) > 1), 'sku_cnt': sku_cnt[c3]})
for c4, c3 in sorted(cat4_parent.items()):
    votes = cat4_votes.get(c4)
    name = CAT4_NAME.get(c4)
    categories.append({'code': c4, 'level': 4, 'parent_code': c3, 'name': name,
                       'name_source': 'МСИ I_PRODUCT (I_CATEGORY_4)' if name else None,
                       'has_name_conflict': bool(votes and len(votes) > 1), 'sku_cnt': sku_cnt[c4]})
for c5, c4 in sorted(cat5_parent.items()):
    categories.append({'code': c5, 'level': 5, 'parent_code': c4, 'name': name5.get(c5),
                       'name_source': src5.get(c5), 'has_name_conflict': conflict5.get(c5, False),
                       'sku_cnt': sku_cnt[c5]})

# ── 7. Товары: агрегаты по городам ─────────────────────────────────────────
agg = collections.defaultdict(lambda: {'shops': 0, 'checks': 0, 'price_num': 0.0, 'cities': 0})
for row in city_rows:
    a = agg[row['sku']]
    a['shops'] += int(row['count_shops'] or 0)
    a['checks'] += int(row['count_check'] or 0)
    a['cities'] += 1
    if row['avg_price'] is not None and row['count_check']:
        a['price_num'] += row['avg_price'] * row['count_check']
for sku, p in products.items():
    a = agg[sku]
    p['id_product'] = EXT2ID.get(sku)
    p['shops_total'] = a['shops']
    p['checks_total'] = a['checks']
    p['cities'] = a['cities']
    p['avg_price'] = round(a['price_num'] / a['checks'], 4) if a['checks'] else None

# ── 8. CSV под \copy ───────────────────────────────────────────────────────
def write_csv(path, fields, rows_):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        w.writeheader()
        for r in rows_:
            w.writerow(r)


CAT_FIELDS = ['code', 'level', 'parent_code', 'name', 'name_source', 'has_name_conflict', 'sku_cnt']
PROD_FIELDS = ['sku', 'sku_code', 'id_product', 'product_name', 'cat5_code', 'cat4_code', 'cat3_code',
               'shops_total', 'checks_total', 'cities', 'avg_price']
CITY_FIELDS = ['sku', 'city'] + [m.lower() for m in METRICS] + ['paired_sku_raw']

write_csv(OUT / 'category.csv', CAT_FIELDS, categories)
write_csv(OUT / 'product.csv', PROD_FIELDS, sorted(products.values(), key=lambda p: p['sku']))
write_csv(OUT / 'product_city.csv', CITY_FIELDS, city_rows)

# ── 9. Excel-справочник ────────────────────────────────────────────────────
name_by_code = {c['code']: c['name'] for c in categories}
out = openpyxl.Workbook()
ws = out.active
ws.title = 'Категории'
ws.append(['Код', 'Уровень', 'Родитель', 'Название', 'Источник названия', 'Разные названия', 'SKU'])
for c in categories:
    ws.append([c['code'], c['level'], c['parent_code'], c['name'], c['name_source'],
               'да' if c['has_name_conflict'] else '', c['sku_cnt']])

ws2 = out.create_sheet('Товары')
ws2.append(['SKU', 'Код SKU', 'ID_PRODUCT', 'Название', 'cat5', 'Категория 5', 'cat4', 'Категория 4',
            'cat3', 'Магазинов', 'Чеков', 'Городов', 'Средняя цена'])
for p in sorted(products.values(), key=lambda x: (-x['checks_total'], x['product_name'])):
    ws2.append([p['sku'], p['sku_code'], p['id_product'], p['product_name'],
                p['cat5_code'], name_by_code.get(p['cat5_code']),
                p['cat4_code'], name_by_code.get(p['cat4_code']), p['cat3_code'],
                p['shops_total'], p['checks_total'], p['cities'], p['avg_price']])

ws3 = out.create_sheet('Товары по городам')
ws3.append(['SKU', 'Название', 'Город'] + METRICS)
pname = {s: p['product_name'] for s, p in products.items()}
for row in city_rows:
    ws3.append([row['sku'], pname[row['sku']], row['city']] + [row[m.lower()] for m in METRICS])
for sheet in (ws, ws2, ws3):
    sheet.freeze_panes = 'A2'
out.save(OUT / 'ассортимент_полный.xlsx')

# ── 10. Отчёт о сборке ──────────────────────────────────────────────────────
lvl = collections.Counter(c['level'] for c in categories)
named = collections.Counter(c['level'] for c in categories if c['name'])
report = {
    'товаров': len(products),
    'строк SKU×город': len(city_rows),
    'городов': sorted({r['city'] for r in city_rows}),
    'категорий': {f'cat{l}': lvl[l] for l in (3, 4, 5)},
    'с названием': {f'cat{l}': named[l] for l in (3, 4, 5)},
    'с расхождением имён': {f'cat{l}': sum(1 for c in categories if c['level'] == l and c['has_name_conflict'])
                            for l in (3, 4, 5)},
    'SKU с внутренним ID_PRODUCT': sum(1 for p in products.values() if p['id_product']),
    'PAIRED_SKU восстановлено из даты': sum(1 for r in city_rows if r['paired_sku_raw']),
}
(OUT / 'build_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
